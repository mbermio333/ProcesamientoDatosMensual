# -*- coding: utf-8 -*-
import os
import json
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image as XLImage
import time

# ------------------ CONFIGURACIÓN GENERAL ------------------
# Cargar configuración desde archivo
CONFIG_FILE = "config_procesamiento.json"

def cargar_configuracion():
    """Cargar configuración desde archivo JSON"""
    config_default = {
        "fm_path": "MedicionesFmCSV",
        "tv_path": "MedicionesTvCSV", 
        "am_path": "MedicionesAmCSV",  # ← NUEVA RUTA PARA AM
        "output_path": "ReportesUnificados",
        "emisoras_por_ciudad": {}
    }
    
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r') as f:
                config = json.load(f)
                # Asegurarse de que exista la clave para emisoras por ciudad
                if "emisoras_por_ciudad" not in config:
                    config["emisoras_por_ciudad"] = {}
                return config
        except:
            return config_default
    
    return config_default

def guardar_configuracion(config):
    """Guardar configuración en archivo JSON"""
    try:
        with open(CONFIG_FILE, 'w') as f:
            json.dump(config, f, indent=4, ensure_ascii=False)
        return True
    except Exception as e:
        print(f"Error al guardar configuración: {e}")
        return False

def extraer_nombres_emisoras(ruta_archivo, tipo):
    """Extraer nombres únicos de emisoras de un archivo CSV con sus frecuencias"""
    try:
        df = pd.read_csv(ruta_archivo, encoding="unicode_escape")
        emisoras_con_frecuencia = []

        columnas_posibles_nombre = ["ESTACION", "Estación", "Station", "STATION", "Nombre de la estación"]
        columna_nombre = None

        for col in columnas_posibles_nombre:
            if col in df.columns:
                columna_nombre = col
                break

        if columna_nombre is None:
            print(f"No se encontró columna de nombre en {ruta_archivo}")
            return []

        if "Frecuencia (Hz)" not in df.columns:
            print(f"No se encontró columna 'Frecuencia (Hz)' en {ruta_archivo}")
            return []

        # Obtener pares únicos estación + frecuencia, sin usar groupby
        grouped = df[[columna_nombre, "Frecuencia (Hz)"]].dropna().drop_duplicates()

        for _, row in grouped.iterrows():
            nombre = str(row[columna_nombre]).strip()
            frecuencia = row["Frecuencia (Hz)"] / 1_000_000  # Convertir a MHz

            if nombre and nombre != "nan" and nombre != "None" and nombre.strip() != "":
                emisoras_con_frecuencia.append({
                    "nombre": nombre,
                    "frecuencia": round(frecuencia, 2),
                    "tipo": tipo
                })

        return emisoras_con_frecuencia

    except Exception as e:
        print(f"Error al extraer nombres de {tipo} desde {ruta_archivo}: {e}")
        return []

def obtener_frecuencias_deseadas(emisoras_config):
    """
    Obtiene las frecuencias deseadas (autorizadas y no autorizadas) excluyendo las con _OBSERVACION
    """
    frecuencias_deseadas = {}
    
    for ciudad, tipos in emisoras_config.items():
        frecuencias_deseadas[ciudad] = {"FM": [], "TV": [], "AM": []}  # ← AGREGAR AM
        
        for tipo in ["FM", "TV", "AM"]:  # ← INCLUIR AM
            if tipo in tipos:
                for emisora in tipos[tipo]:
                    # Excluir frecuencias con _OBSERVACION
                    if "_OBSERVACION" not in emisora.get("nombre", "").upper():
                        # Usar string para la frecuencia para consistencia
                        frecuencia_str = str(emisora.get("frecuencia", ""))
                        if frecuencia_str and frecuencia_str != "nan":
                            frecuencias_deseadas[ciudad][tipo].append({
                                "nombre": emisora.get("nombre", ""),
                                "frecuencia": frecuencia_str,
                                "tipo": tipo
                            })
    
    return frecuencias_deseadas

def extraer_frecuencias_csv(ruta_archivo, tipo):
    """
    Extrae todas las frecuencias de un archivo CSV (incluyendo nombres en blanco)
    """
    try:
        df = pd.read_csv(ruta_archivo, encoding="unicode_escape")
        frecuencias_csv = []

        columnas_posibles_nombre = ["ESTACION", "Estación", "Station", "STATION", "Nombre de la estación"]
        columna_nombre = None

        for col in columnas_posibles_nombre:
            if col in df.columns:
                columna_nombre = col
                break

        if columna_nombre is None:
            print(f"No se encontró columna de nombre en {ruta_archivo}")
            return []

        if "Frecuencia (Hz)" not in df.columns:
            print(f"No se encontró columna 'Frecuencia (Hz)' en {ruta_archivo}")
            return []

        # Obtener todos los pares únicos estación + frecuencia
        grouped = df[[columna_nombre, "Frecuencia (Hz)"]].drop_duplicates()

        for _, row in grouped.iterrows():
            nombre_csv = str(row[columna_nombre]).strip() if pd.notna(row[columna_nombre]) else ""
            frecuencia = row["Frecuencia (Hz)"] / 1_000_000  # Convertir a MHz
            frecuencia_str = str(round(frecuencia, 2))

            # Incluir todas las frecuencias, incluso las con nombres en blanco
            frecuencias_csv.append({
                "nombre": nombre_csv,
                "frecuencia": frecuencia_str,
                "tipo": tipo
            })

        return frecuencias_csv

    except Exception as e:
        print(f"Error al extraer frecuencias de {tipo} desde {ruta_archivo}: {e}")
        return []

def cotejar_y_actualizar_frecuencias(config, archivos_fm, archivos_tv, archivos_am, callback_log=None):
    """
    Coteja las frecuencias entre config.json y los archivos CSV, pero NO actualiza config.json
    Retorna la lista final de frecuencias a procesar respetando los nombres de config.json
    """
    frecuencias_a_procesar = {}
    
    # Obtener frecuencias deseadas actuales (excluyendo _OBSERVACION)
    frecuencias_deseadas = obtener_frecuencias_deseadas(config.get("emisoras_por_ciudad", {}))
    
    # Procesar cada ciudad NORMALIZANDO el nombre
    for ciudad_original in config.get("emisoras_por_ciudad", {}).keys():
        ciudad_normalizada = normalizar_nombre_ciudad(ciudad_original)
        
        # Buscar la ciudad normalizada en los archivos
        ciudad_en_archivos = None
        for archivo_ciudad in archivos_fm.keys():
            if normalizar_nombre_ciudad(archivo_ciudad) == ciudad_normalizada:
                ciudad_en_archivos = archivo_ciudad
                break
        
        if not ciudad_en_archivos:
            # Intentar en archivos TV si no se encuentra en FM
            for archivo_ciudad in archivos_tv.keys():
                if normalizar_nombre_ciudad(archivo_ciudad) == ciudad_normalizada:
                    ciudad_en_archivos = archivo_ciudad
                    break
        
        if not ciudad_en_archivos:
            # Intentar en archivos AM si no se encuentra en FM o TV
            for archivo_ciudad in archivos_am.keys():
                if normalizar_nombre_ciudad(archivo_ciudad) == ciudad_normalizada:
                    ciudad_en_archivos = archivo_ciudad
                    break
        
        if not ciudad_en_archivos:
            if callback_log:
                callback_log(f"❌ Ciudad {ciudad_original} (normalizada: {ciudad_normalizada}) no encontrada en archivos CSV")
            continue
            
        frecuencias_a_procesar[ciudad_en_archivos] = {"FM": [], "TV": [], "AM": []}
        
        # Procesar FM
        if ciudad_en_archivos in archivos_fm:
            frecuencias_csv_fm = extraer_frecuencias_csv(archivos_fm[ciudad_en_archivos], "FM")
            
            for freq_deseada in frecuencias_deseadas.get(ciudad_original, {}).get("FM", []):
                freq_str_deseada = freq_deseada["frecuencia"]
                nombre_deseado = freq_deseada["nombre"]
                
                # Buscar esta frecuencia en el CSV
                frecuencia_encontrada = None
                for freq_csv in frecuencias_csv_fm:
                    if freq_csv["frecuencia"] == freq_str_deseada:
                        frecuencia_encontrada = freq_csv
                        break
                
                if frecuencia_encontrada:
                    nombre_csv = frecuencia_encontrada["nombre"]
                    
                    # Caso 1: Nombre en CSV está en blanco, usar el de config
                    if not nombre_csv or nombre_csv == "nan" or nombre_csv == "None":
                        frecuencia_a_procesar = {
                            "nombre": nombre_deseado,
                            "frecuencia": freq_str_deseada,
                            "tipo": "FM"
                        }
                        frecuencias_a_procesar[ciudad_en_archivos]["FM"].append(frecuencia_a_procesar)
                        
                    # Caso 2: Nombre en CSV es diferente al de config - RESPETAR CONFIG (NO ACTUALIZAR)
                    elif nombre_csv != nombre_deseado:
                        if callback_log:
                            callback_log(f"⚠️  Diferencia encontrada en {ciudad_en_archivos} FM {freq_str_deseada}: Config='{nombre_deseado}' vs CSV='{nombre_csv}' - RESPETANDO CONFIG")
                        
                        # USAR SIEMPRE EL NOMBRE DE CONFIG (no actualizar el archivo)
                        frecuencia_a_procesar = {
                            "nombre": nombre_deseado,  # ← Usar siempre el nombre de config
                            "frecuencia": freq_str_deseada,
                            "tipo": "FM"
                        }
                        frecuencias_a_procesar[ciudad_en_archivos]["FM"].append(frecuencia_a_procesar)
                    
                    # Caso 3: Nombres iguales - procesar normalmente
                    else:
                        frecuencia_a_procesar = {
                            "nombre": nombre_deseado,
                            "frecuencia": freq_str_deseada,
                            "tipo": "FM"
                        }
                        frecuencias_a_procesar[ciudad_en_archivos]["FM"].append(frecuencia_a_procesar)
                
                else:
                    # Frecuencia deseada no encontrada en CSV
                    if callback_log:
                        callback_log(f"❌ Frecuencia FM {freq_str_deseada} ({nombre_deseado}) no encontrada en CSV de {ciudad_en_archivos}")
        
        # Procesar TV (misma lógica que FM)
        if ciudad_en_archivos in archivos_tv:
            frecuencias_csv_tv = extraer_frecuencias_csv(archivos_tv[ciudad_en_archivos], "TV")
            
            for freq_deseada in frecuencias_deseadas.get(ciudad_original, {}).get("TV", []):
                freq_str_deseada = freq_deseada["frecuencia"]
                nombre_deseado = freq_deseada["nombre"]
                
                # Buscar esta frecuencia en el CSV
                frecuencia_encontrada = None
                for freq_csv in frecuencias_csv_tv:
                    if freq_csv["frecuencia"] == freq_str_deseada:
                        frecuencia_encontrada = freq_csv
                        break
                
                if frecuencia_encontrada:
                    nombre_csv = frecuencia_encontrada["nombre"]
                    
                    # Caso 1: Nombre en CSV está en blanco, usar el de config
                    if not nombre_csv or nombre_csv == "nan" or nombre_csv == "None":
                        frecuencia_a_procesar = {
                            "nombre": nombre_deseado,
                            "frecuencia": freq_str_deseada,
                            "tipo": "TV"
                        }
                        frecuencias_a_procesar[ciudad_en_archivos]["TV"].append(frecuencia_a_procesar)
                        
                    # Caso 2: Nombre en CSV es diferente al de config - RESPETAR CONFIG (NO ACTUALIZAR)
                    elif nombre_csv != nombre_deseado:
                        if callback_log:
                            callback_log(f"⚠️  Diferencia encontrada en {ciudad_en_archivos} TV {freq_str_deseada}: Config='{nombre_deseado}' vs CSV='{nombre_csv}' - RESPETANDO CONFIG")
                        
                        # USAR SIEMPRE EL NOMBRE DE CONFIG (no actualizar el archivo)
                        frecuencia_a_procesar = {
                            "nombre": nombre_deseado,  # ← Usar siempre el nombre de config
                            "frecuencia": freq_str_deseada,
                            "tipo": "TV"
                        }
                        frecuencias_a_procesar[ciudad_en_archivos]["TV"].append(frecuencia_a_procesar)
                    
                    # Caso 3: Nombres iguales - procesar normalmente
                    else:
                        frecuencia_a_procesar = {
                            "nombre": nombre_deseado,
                            "frecuencia": freq_str_deseada,
                            "tipo": "TV"
                        }
                        frecuencias_a_procesar[ciudad_en_archivos]["TV"].append(frecuencia_a_procesar)
                
                else:
                    # Frecuencia deseada no encontrada en CSV
                    if callback_log:
                        callback_log(f"❌ Frecuencia TV {freq_str_deseada} ({nombre_deseado}) no encontrada en CSV de {ciudad_en_archivos}")

        # Procesar AM
        if ciudad_en_archivos in archivos_am:
            frecuencias_csv_am = extraer_frecuencias_csv(archivos_am[ciudad_en_archivos], "AM")
            
            for freq_deseada in frecuencias_deseadas.get(ciudad_original, {}).get("AM", []):
                freq_str_deseada = freq_deseada["frecuencia"]
                nombre_deseado = freq_deseada["nombre"]
                
                # Buscar esta frecuencia en el CSV
                frecuencia_encontrada = None
                for freq_csv in frecuencias_csv_am:
                    if freq_csv["frecuencia"] == freq_str_deseada:
                        frecuencia_encontrada = freq_csv
                        break
                
                if frecuencia_encontrada:
                    nombre_csv = frecuencia_encontrada["nombre"]
                    
                    # Caso 1: Nombre en CSV está en blanco, usar el de config
                    if not nombre_csv or nombre_csv == "nan" or nombre_csv == "None":
                        frecuencia_a_procesar = {
                            "nombre": nombre_deseado,
                            "frecuencia": freq_str_deseada,
                            "tipo": "AM"
                        }
                        frecuencias_a_procesar[ciudad_en_archivos]["AM"].append(frecuencia_a_procesar)
                        
                    # Caso 2: Nombre en CSV es diferente al de config - RESPETAR CONFIG (NO ACTUALIZAR)
                    elif nombre_csv != nombre_deseado:
                        if callback_log:
                            callback_log(f"⚠️  Diferencia encontrada en {ciudad_en_archivos} AM {freq_str_deseada}: Config='{nombre_deseado}' vs CSV='{nombre_csv}' - RESPETANDO CONFIG")
                        
                        # USAR SIEMPRE EL NOMBRE DE CONFIG (no actualizar el archivo)
                        frecuencia_a_procesar = {
                            "nombre": nombre_deseado,  # ← Usar siempre el nombre de config
                            "frecuencia": freq_str_deseada,
                            "tipo": "AM"
                        }
                        frecuencias_a_procesar[ciudad_en_archivos]["AM"].append(frecuencia_a_procesar)
                    
                    # Caso 3: Nombres iguales - procesar normalmente
                    else:
                        frecuencia_a_procesar = {
                            "nombre": nombre_deseado,
                            "frecuencia": freq_str_deseada,
                            "tipo": "AM"
                        }
                        frecuencias_a_procesar[ciudad_en_archivos]["AM"].append(frecuencia_a_procesar)
                
                else:
                    # Frecuencia deseada no encontrada en CSV
                    if callback_log:
                        callback_log(f"❌ Frecuencia AM {freq_str_deseada} ({nombre_deseado}) no encontrada en CSV de {ciudad_en_archivos}")

    # ELIMINADO: No guardar configuración automáticamente
    # El archivo config.json permanece sin cambios
    
    return frecuencias_a_procesar


def limpiar_configuracion_duplicados(config):
    """Une entradas duplicadas de ciudades en config.json"""
    ciudades_normalizadas = {}
    
    for ciudad, datos in config.get("emisoras_por_ciudad", {}).items():
        ciudad_norm = normalizar_nombre_ciudad(ciudad)
        
        if ciudad_norm not in ciudades_normalizadas:
            ciudades_normalizadas[ciudad_norm] = datos
        else:
            # Unir datos duplicados para FM, TV y AM
            for tipo in ["FM", "TV", "AM"]:  # ← AGREGAR AM
                if tipo in datos:
                    if tipo not in ciudades_normalizadas[ciudad_norm]:
                        ciudades_normalizadas[ciudad_norm][tipo] = []
                    
                    # Evitar duplicados por frecuencia
                    frecuencias_existentes = {str(e["frecuencia"]) for e in ciudades_normalizadas[ciudad_norm][tipo]}
                    for emisora in datos[tipo]:
                        if str(emisora["frecuencia"]) not in frecuencias_existentes:
                            ciudades_normalizadas[ciudad_norm][tipo].append(emisora)
                            frecuencias_existentes.add(str(emisora["frecuencia"]))
    
    config["emisoras_por_ciudad"] = ciudades_normalizadas
    return config


def filtrar_dataframe_por_frecuencias(df, frecuencias_procesar, tipo, callback_log=None):
    """
    Filtra el DataFrame para incluir solo las frecuencias deseadas
    """
    try:
        # Convertir frecuencia a string para comparación consistente
        df["Frecuencia (MHz)"] = (df["Frecuencia (Hz)"] / 1_000_000).round(2)
        df["Frecuencia_Str"] = df["Frecuencia (MHz)"].astype(str)
        
        # Obtener lista de frecuencias a incluir
        frecuencias_incluir = [f["frecuencia"] for f in frecuencias_procesar]
        
        # Filtrar DataFrame
        df_filtrado = df[df["Frecuencia_Str"].isin(frecuencias_incluir)].copy()
        
        # Asignar nombres correctos desde la configuración
        for frecuencia_info in frecuencias_procesar:
            freq_str = frecuencia_info["frecuencia"]
            nombre_correcto = frecuencia_info["nombre"]
            
            # Actualizar nombres en el DataFrame
            mask = df_filtrado["Frecuencia_Str"] == freq_str
            if mask.any():
                # Usar la columna correcta para el nombre
                columna_nombre = "ESTACION" if "ESTACION" in df_filtrado.columns else "Nombre de la estación"
                if columna_nombre in df_filtrado.columns:
                    df_filtrado.loc[mask, columna_nombre] = nombre_correcto
        
        # Eliminar columna temporal
        df_filtrado = df_filtrado.drop(columns=["Frecuencia_Str"])
        
        if callback_log:
            callback_log(f"✅ DataFrame {tipo} filtrado: {len(df_filtrado)} registros de {len(frecuencias_incluir)} frecuencias")
        
        return df_filtrado
        
    except Exception as e:
        if callback_log:
            callback_log(f"❌ Error filtrando DataFrame {tipo}: {str(e)}")
        return df


# Cargar configuración al inicio
# Agregar ruta AM a las variables globales
config = cargar_configuracion()
ruta_fm = config.get("fm_path", "MedicionesFmCSV")
ruta_tv = config.get("tv_path", "MedicionesTvCSV")
ruta_am = config.get("am_path", "MedicionesAmCSV")  # ← NUEVA VARIABLE
ruta_salida = config.get("output_path", "ReportesUnificados")
ruta_imagenes = "Img"
fecha_actual = datetime.now().strftime("%d/%m/%Y")
# Llamar esta función después de cargar la configuración



# ------------------ FUNCIONES AUXILIARES ------------------

def inicializar_directorios():
    """Crear los directorios necesarios si no existen"""
    os.makedirs(ruta_salida, exist_ok=True)
    os.makedirs(ruta_fm, exist_ok=True)
    os.makedirs(ruta_tv, exist_ok=True)
    os.makedirs(ruta_am, exist_ok=True)  # ← NUEVO DIRECTORIO
    os.makedirs(ruta_imagenes, exist_ok=True)

def combinar_observaciones_solo_en_tv(ws, fila_inicio_tabla, fila_fin_tabla, num_columnas, es_tv=True):
    """
    Combina horizontalmente la última y penúltima columna para cada fila SOLO si es TV.
    """
    if es_tv:
        for fila in range(fila_inicio_tabla + 1, fila_fin_tabla + 1):
            ws.merge_cells(start_row=fila, start_column=num_columnas - 1, end_row=fila, end_column=num_columnas)
        
        # También combinar el encabezado (la celda con el nombre del campo)
        ws.merge_cells(start_row=fila_inicio_tabla, start_column=num_columnas - 1, end_row=fila_inicio_tabla, end_column=num_columnas)
        celda = ws.cell(row=fila_inicio_tabla, column=num_columnas - 1)
        celda.value = "OBSERVACIONES"
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row=fila_inicio_tabla, column=num_columnas).value = None

def formatear_hoja(ws, fila_inicio, encabezado_lineas=None):
    """Formatea una hoja de cálculo con bordes y estilos"""
    if encabezado_lineas:
        num_columnas = ws.max_column
        ws.insert_rows(fila_inicio, amount=len(encabezado_lineas))

        for i, texto in enumerate(encabezado_lineas):
            fila = fila_inicio + i
            celda = ws.cell(row=fila, column=1, value=texto)
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=num_columnas)
            celda.alignment = Alignment(horizontal="center", vertical="center")
            celda.font = Font(bold=True, size=12)
    
    # Bordes
    inicio_fila_tabla = fila_inicio + (len(encabezado_lineas) if encabezado_lineas else 0)
    fin_fila_tabla = ws.max_row
    thin = Side(border_style="thin")
    borde_grueso = Side(border_style="medium")

    for row in range(inicio_fila_tabla, fin_fila_tabla + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if row == inicio_fila_tabla:
                cell.border = Border(top=borde_grueso, bottom=thin,
                                     left=borde_grueso if col == 1 else thin,
                                     right=borde_grueso if col == ws.max_column else thin)
            elif row == fin_fila_tabla:
                cell.border = Border(top=thin, bottom=borde_grueso,
                                     left=borde_grueso if col == 1 else thin,
                                     right=borde_grueso if col == ws.max_column else thin)
            elif col == 1:
                cell.border = Border(left=borde_grueso, top=thin, bottom=thin, right=thin)
            elif col == ws.max_column:
                cell.border = Border(right=borde_grueso, top=thin, bottom=thin, left=thin)

            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Ajustar anchos de columnas
    for col in range(1, ws.max_column + 1):
        celda = ws.cell(row=inicio_fila_tabla, column=col)
        if celda.value == "Promedio(dBuV/m)":
            celda.value = "Promedio\n(dBuV/m)"
            ws.column_dimensions[get_column_letter(col)].width = 15
        elif celda.value == "Ancho de Banda (KHz)":
            celda.value = "Ancho de Banda\n(KHz)"
        elif celda.value == "Medición Manual":
            celda.value = "Medición Manual\nAB(KHz) o NIVEL\n(dBµV/m)"
            ws.column_dimensions[get_column_letter(col)].width = 23
        celda.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    # Ajustar anchos del resto
    for i in range(1, ws.max_column + 1):
        if ws.column_dimensions[get_column_letter(i)].width in [15, 23]:
            continue
        col_letter = get_column_letter(i)
        max_length = 0
        for cell in ws[col_letter]:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = max_length + 2

    if encabezado_lineas:
        ws.row_dimensions[inicio_fila_tabla].height = 45

def insertar_imagenes(ws, fila_destino, tipo):
    logo_izquierda = os.path.join(ruta_imagenes, "ARCOTEL.png")
    logo_derecha = os.path.join(ruta_imagenes, "nEcuador.png")

    if os.path.exists(logo_izquierda):
        img_left = XLImage(logo_izquierda)
        img_left.width = 500
        img_left.height = 100
        ws.add_image(img_left, f"A{fila_destino}")

    if os.path.exists(logo_derecha):
        img_right = XLImage(logo_derecha)
        img_right.width = 260
        img_right.height = 120
        ws.add_image(img_right, f"AJ{fila_destino}")


def normalizar_nombre_ciudad(nombre):
    """Normaliza el nombre de la ciudad para consistencia"""
    if not nombre or not isinstance(nombre, str):
        return ""
    
    nombre = nombre.lower().strip()
    
    # Manejar todas las variantes de "cañar" de manera consistente
    ##if nombre in ["cañar", "cañar", "canar", "caã±ar", "tambo"]:
    #    return "cañar"  # ← DEVOLVER SIEMPRE LA MISMA CLAVE
    
    # Para otras ciudades, devolver en minúsculas para consistencia
    mapeo_ciudades = {
        "zamora": "zamora",
        "loja": "loja", 
        "macas": "macas",
        "tambo":"tambo",
        "machala": "machala",
        "cuenca": "cuenca"
        
    }
    
    return mapeo_ciudades.get(nombre, nombre.lower())

def limpiar_configuracion_duplicados(config):
    """Une entradas duplicadas de ciudades en config.json"""
    ciudades_normalizadas = {}
    
    for ciudad, datos in config.get("emisoras_por_ciudad", {}).items():
        ciudad_norm = normalizar_nombre_ciudad(ciudad)
        
        if ciudad_norm not in ciudades_normalizadas:
            ciudades_normalizadas[ciudad_norm] = datos
        else:
            # Unir datos duplicados
            for tipo in ["FM", "TV"]:
                if tipo in datos:
                    if tipo not in ciudades_normalizadas[ciudad_norm]:
                        ciudades_normalizadas[ciudad_norm][tipo] = []
                    
                    # Evitar duplicados por frecuencia
                    frecuencias_existentes = {str(e["frecuencia"]) for e in ciudades_normalizadas[ciudad_norm][tipo]}
                    for emisora in datos[tipo]:
                        if str(emisora["frecuencia"]) not in frecuencias_existentes:
                            ciudades_normalizadas[ciudad_norm][tipo].append(emisora)
                            frecuencias_existentes.add(str(emisora["frecuencia"]))
    
    config["emisoras_por_ciudad"] = ciudades_normalizadas
    return config
# ------------------ FUNCIÓN PARA COLOREAR CELDAS ------------------

def colorear_celdas_por_valor(ws, fila_inicio, fila_fin, tipo, col_names):
    from openpyxl.styles import PatternFill

    # DEFINIR TODOS LOS COLORES AL INICIO DE LA FUNCIÓN
    rojo = PatternFill(start_color="FFFC4A2C", end_color="FFFC4A2C", fill_type="solid")
    amarillo = PatternFill(start_color="FFFFFE9F", end_color="FFFFFE9F", fill_type="solid")
    verde = PatternFill(start_color="FFCDFECE", end_color="FFCDFECE", fill_type="solid")
    rosa = PatternFill(start_color="FFFD9BCB", end_color="FFFD9BCB", fill_type="solid")

    encabezados = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=fila_inicio - 1, column=col).value
        if isinstance(val, str):
            val_normalizado = val.strip().replace('\n', ' ')
            encabezados[val_normalizado] = col
        elif isinstance(val, (int, float)):
            encabezados[str(val)] = col

    def obtener_banda(freq):
        if (54 <= freq <= 72) or (76 <= freq <= 88):
            return 'I'
        elif 174 <= freq <= 216:
            return 'III'
        elif (470 <= freq <= 488) or (512 <= freq <= 608):
            return 'IV'
        elif 614 <= freq <= 698:
            return 'V'
        else:
            return None

    def get_umbral_color(banda, valor):
        if banda == 'I':
            if valor < 47:
                return rosa
            elif 47 <= valor < 68:
                return amarillo
            else:
                return verde
        elif banda == 'III':
            if valor < 56:
                return rosa
            elif 56 <= valor < 71:
                return amarillo
            else:
                return verde
        elif banda in ['IV', 'V']:
            if valor < 64:
                return rosa
            elif 64 <= valor < 74:
                return amarillo
            else:
                return verde
        else:
            return None

    for fila in range(fila_inicio, fila_fin + 1):
        frecuencia_col = None
        for key in encabezados:
            if "frecuencia" in key.lower():
                frecuencia_col = encabezados[key]
                break

        frecuencia = ws.cell(row=fila, column=frecuencia_col).value
        try:
            frecuencia = float(frecuencia)
            banda = obtener_banda(frecuencia)
        except:
            banda = None

        for nombre_col in col_names:
            nombre_normalizado = str(nombre_col).strip().replace('\n', ' ')
            if nombre_normalizado not in encabezados:
                continue

            col = encabezados[nombre_normalizado]
            celda = ws.cell(row=fila, column=col)
            valor = celda.value

            try:
                valor = float(valor)
            except (ValueError, TypeError):
                continue

            if tipo == "FM":
                if nombre_normalizado == "Promedio (dBuV/m)":
                    if 0 <= valor <= 30:
                        celda.fill = rojo
                    elif 30 < valor < 54:
                        celda.fill = amarillo
                    elif valor >= 54:
                        celda.fill = verde
                elif nombre_normalizado == "Ancho de Banda (KHz)":
                    if valor <= 220:
                        celda.fill = verde
                    elif valor > 220:
                        celda.fill = rojo
                         # Escribir mensaje en la columna "OBSERVACIONES"
                        col_obs = encabezados.get("OBSERVACIONES")
                        if col_obs:
                            celda_obs = ws.cell(row=fila, column=col_obs)
                            mensaje = "Opera con ancho de banda mayor\na lo autorizado\n(medición automática)"
                            celda_obs.value = mensaje
                            celda_obs.alignment = Alignment(wrap_text=True)
                elif nombre_normalizado in [str(d) for d in range(1, 32)]:
                    if 0 <= valor <= 30:
                        celda.fill = rosa
                    elif 30 < valor < 54:
                        celda.fill = amarillo
                    elif valor >= 54:
                        celda.fill = verde

            elif tipo == "TV":
                if banda:
                    if nombre_normalizado == "Promedio (dBuV/m)" or nombre_normalizado in [str(d) for d in range(1, 32)]:
                        frecuencia_col = encabezados.get("Frecuencia (MHz)")
                        observaciones_col = encabezados.get("OBSERVACIONES")
                        # Dentro de if nombre_normalizado == "Promedio (dBuV/m)":
                        frecuencia = ws.cell(row=fila, column=frecuencia_col).value if frecuencia_col else None

                        if frecuencia:
                            banda = obtener_banda(frecuencia)
                            color = get_umbral_color(banda, valor)

                            if color:
                                celda.fill = color
                                # Agrega mensaje si es rosa (menor al mínimo por banda)
                                if color == rosa and observaciones_col:
                                    obs_cell = ws.cell(row=fila, column=observaciones_col)
                                    existing = obs_cell.value or ""
                                    mensaje = "Niveles por debajo del borde del área de cobertura principal y secundaria"
                                    
                                    if mensaje not in existing:
                                        nuevo_texto = (existing + "\n" + mensaje).strip()
                                        obs_cell.value = nuevo_texto
                                        obs_cell.alignment = Alignment(wrap_text=True, vertical="center")
                                        
                                        # Estimar número de líneas para ajustar alto de fila
                                        num_lineas = nuevo_texto.count("\n") + 1
                                        altura = num_lineas * 15  # Puedes ajustar este valor según cómo se vea
                                        ws.row_dimensions[fila].height = altura

            # NUEVO: Reglas para AM
            elif tipo == "AM":
                if nombre_normalizado == "Promedio (dBuV/m)":
                    if valor < 62:
                        celda.fill = rosa
                    else:
                        celda.fill = verde
                elif nombre_normalizado in [str(d) for d in range(1, 32)]:
                    if valor < 62:
                        celda.fill = rosa
                    else:
                        celda.fill = verde



def encontrar_columnas_numericas(ws, fila_encabezados):
    """
    Encuentra automáticamente las columnas que contienen valores numéricos
    basándose en los encabezados de días (1 al 31)
    """
    columnas_numericas = []
    
    for col in range(1, ws.max_column + 1):
        celda = ws.cell(row=fila_encabezados, column=col)
        if celda.value and str(celda.value).isdigit():
            try:
                dia = int(celda.value)
                if 1 <= dia <= 31:
                    columnas_numericas.append(col)
            except ValueError:
                pass
    
    if columnas_numericas:
        return min(columnas_numericas), max(columnas_numericas)
    else:
        # Valores por defecto si no encuentra días
        return 3, 33

def obtener_base(nombre_archivo):
    """Obtiene el nombre base del archivo y filtra 'global'"""
    base = nombre_archivo.split("_")[0].lower().strip()
    # Filtrar "global" y variantes
    if base in ["global", "global", "generico", "general"]:
        return None
    return base

# ------------------ FUNCIONES PARA CREAR HOJAS ADICIONALES ------------------

def crear_hoja_manual(wb, nombre_hoja):
    """Crea una hoja en blanco con el nombre especificado"""
    if nombre_hoja not in wb.sheetnames:
        ws = wb.create_sheet(nombre_hoja)
        # Agregar encabezados básicos
        if "FM" in nombre_hoja:
            encabezados = ["ESTACION", "Frecuencia (MHz)", "Medición Manual AB(KHz) o NIVEL (dBµV/m)", "OBSERVACIONES"]
        else:
            encabezados = ["ESTACION", "Frecuencia (MHz)", "Medición Manual AB(KHz) o NIVEL (dBµV/m)", "OBSERVACIONES"]
        
        for col, encabezado in enumerate(encabezados, 1):
            ws.cell(row=1, column=col, value=encabezado)
        
        # Aplicar formato básico
        for col in range(1, len(encabezados) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
            celda = ws.cell(row=1, column=col)
            celda.font = Font(bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center")
        
        # Aplicar bordes a la fila de encabezados
        thin = Side(border_style="thin")
        for col in range(1, len(encabezados) + 1):
            ws.cell(row=1, column=col).border = Border(top=thin, bottom=thin, left=thin, right=thin)

def crear_hoja_observaciones(wb, datos_fm, datos_tv, datos_am=None):
    """Crea la hoja de observaciones con los datos de FM, TV y AM ordenados por frecuencia"""
    if "Observaciones" in wb.sheetnames:
        ws_obs = wb["Observaciones"]
    else:
        ws_obs = wb.create_sheet("Observaciones")
    
    # Limpiar hoja existente
    ws_obs.delete_rows(1, ws_obs.max_row)
    
    fila_actual = 1
    thin = Side(border_style="thin")
    borde_grueso = Side(border_style="medium")
    
    # DEFINIR COLORES AL INICIO
    from openpyxl.styles import PatternFill
    rojo = PatternFill(start_color="FFFC4A2C", end_color="FFFC4A2C", fill_type="solid")
    amarillo = PatternFill(start_color="FFFFFE9F", end_color="FFFFFE9F", fill_type="solid")
    verde = PatternFill(start_color="FFCDFECE", end_color="FFCDFECE", fill_type="solid")
    rosa = PatternFill(start_color="FFFD9BCB", end_color="FFFD9BCB", fill_type="solid")
    
    # Definir anchos específicos para columnas
    anchos_especificos = {
        "ESTACION": 25,
        "Frecuencia (MHz)": 18,
        "Promedio(dBuV/m)": 18,
        "Ancho de Banda (KHz)": 20,
        "Medición Manual AB(KHz) o NIVEL (dBµV/m)": 20,
        "OBSERVACIONES": 35
    }
    
    # Agregar datos de FM (ordenados por frecuencia)
    if datos_fm is not None and not datos_fm.empty:
        # Ordenar datos FM por frecuencia (de menor a mayor)
        datos_fm_ordenados = datos_fm.sort_values(by="Frecuencia (MHz)")
        
        # Encabezado para FM
        ws_obs.cell(row=fila_actual, column=1, value="FM")
        ws_obs.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=6)
        celda = ws_obs.cell(row=fila_actual, column=1)
        celda.font = Font(bold=True, size=14)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        fila_actual += 1
        
        # Encabezados de columnas para FM
        encabezados_fm = ["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", "Ancho de Banda (KHz)", 
                         "Medición Manual AB(KHz)\no NIVEL (dBµV/m)", "OBSERVACIONES"]
        
        for col, encabezado in enumerate(encabezados_fm, 1):
            ws_obs.cell(row=fila_actual, column=col, value=encabezado)
            celda = ws_obs.cell(row=fila_actual, column=col)
            celda.font = Font(bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            # Borde grueso para encabezados
            celda.border = Border(top=borde_grueso, bottom=thin, left=borde_grueso if col == 1 else thin, 
                                 right=borde_grueso if col == len(encabezados_fm) else thin)
            
            # Aplicar ancho específico si existe en el diccionario
            if encabezado in anchos_especificos:
                ws_obs.column_dimensions[get_column_letter(col)].width = anchos_especificos[encabezado]
        
        fila_inicio_fm = fila_actual
        fila_actual += 1
        
        # Datos de FM ordenados
        for _, row in datos_fm_ordenados.iterrows():
            for col_idx, col_name in enumerate(["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", 
                                              "Ancho de Banda (KHz)", "Medición Manual", "OBSERVACIONES"], 1):
                ws_obs.cell(row=fila_actual, column=col_idx, value=row[col_name])
                # Aplicar alineación centrada a todas las celdas
                ws_obs.cell(row=fila_actual, column=col_idx).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            
            # Aplicar bordes a cada fila de datos
            for col in range(1, 7):
                celda = ws_obs.cell(row=fila_actual, column=col)
                celda.border = Border(top=thin, bottom=thin, 
                                     left=borde_grueso if col == 1 else thin,
                                     right=borde_grueso if col == 6 else thin)
            
            fila_actual += 1
        
        # Borde inferior grueso para la tabla FM
        for col in range(1, 7):
            celda = ws_obs.cell(row=fila_actual-1, column=col)
            current_border = celda.border
            celda.border = Border(top=current_border.top, bottom=borde_grueso,
                                 left=current_border.left, right=current_border.right)
        
        fila_fin_fm = fila_actual - 1
        fila_actual += 2  # Espacio de 2 filas entre tablas
    
    # Agregar datos de TV (ordenados por frecuencia)
    if datos_tv is not None and not datos_tv.empty:
        # Ordenar datos TV por frecuencia (de menor a mayor)
        datos_tv_ordenados = datos_tv.sort_values(by="Frecuencia (MHz)")
        
        # Encabezado para TV
        ws_obs.cell(row=fila_actual, column=1, value="TV")
        ws_obs.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=5)
        celda = ws_obs.cell(row=fila_actual, column=1)
        celda.font = Font(bold=True, size=14)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        fila_actual += 1
        
        # Encabezados de columnas para TV
        encabezados_tv = ["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", 
                         "Medición Manual AB(KHz)\no NIVEL (dBµV/m)", "OBSERVACIONES"]
        
        for col, encabezado in enumerate(encabezados_tv, 1):
            ws_obs.cell(row=fila_actual, column=col, value=encabezado)
            celda = ws_obs.cell(row=fila_actual, column=col)
            celda.font = Font(bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            # Borde grueso para encabezados
            celda.border = Border(top=borde_grueso, bottom=thin, left=borde_grueso if col == 1 else thin, 
                                 right=borde_grueso if col == len(encabezados_tv) else thin)
            
            # Aplicar ancho específico si existe en el diccionario
            if encabezado in anchos_especificos:
                ws_obs.column_dimensions[get_column_letter(col)].width = anchos_especificos[encabezado]
        
        fila_inicio_tv = fila_actual
        fila_actual += 1
        
        # Datos de TV ordenados
        for _, row in datos_tv_ordenados.iterrows():
            for col_idx, col_name in enumerate(["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", 
                                              "Medición Manual", "OBSERVACIONES"], 1):
                ws_obs.cell(row=fila_actual, column=col_idx, value=row[col_name])
                # Aplicar alineación centrada a todas las celdas
                ws_obs.cell(row=fila_actual, column=col_idx).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            
            # Aplicar bordes a cada fila de datos
            for col in range(1, 6):
                celda = ws_obs.cell(row=fila_actual, column=col)
                celda.border = Border(top=thin, bottom=thin, 
                                     left=borde_grueso if col == 1 else thin,
                                     right=borde_grueso if col == 5 else thin)
            
            fila_actual += 1
        
        # Borde inferior grueso para la tabla TV
        for col in range(1, 6):
            celda = ws_obs.cell(row=fila_actual-1, column=col)
            current_border = celda.border
            celda.border = Border(top=current_border.top, bottom=borde_grueso,
                                 left=current_border.left, right=current_border.right)
        
        fila_fin_tv = fila_actual - 1

    # Después de procesar TV, agregar AM
    if datos_am is not None and not datos_am.empty:
        # Ordenar datos AM por frecuencia (de menor a mayor)
        datos_am_ordenados = datos_am.sort_values(by="Frecuencia (MHz)")
        
        # Encabezado para AM
        ws_obs.cell(row=fila_actual, column=1, value="AM")
        ws_obs.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=5)
        celda = ws_obs.cell(row=fila_actual, column=1)
        celda.font = Font(bold=True, size=14)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        fila_actual += 1
        
        # Encabezados de columnas para AM (mismo formato que TV)
        encabezados_am = ["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", 
                         "Medición Manual AB(KHz)\no NIVEL (dBµV/m)", "OBSERVACIONES"]
        
        for col, encabezado in enumerate(encabezados_am, 1):
            ws_obs.cell(row=fila_actual, column=col, value=encabezado)
            celda = ws_obs.cell(row=fila_actual, column=col)
            celda.font = Font(bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            # Borde grueso para encabezados
            celda.border = Border(top=borde_grueso, bottom=thin, left=borde_grueso if col == 1 else thin, 
                                 right=borde_grueso if col == len(encabezados_am) else thin)
            
            # Aplicar ancho específico si existe en el diccionario
            if encabezado in anchos_especificos:
                ws_obs.column_dimensions[get_column_letter(col)].width = anchos_especificos[encabezado]
        
        fila_inicio_am = fila_actual
        fila_actual += 1
        
        # Datos de AM ordenados
        for _, row in datos_am_ordenados.iterrows():
            for col_idx, col_name in enumerate(["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", 
                                              "Medición Manual", "OBSERVACIONES"], 1):
                ws_obs.cell(row=fila_actual, column=col_idx, value=row[col_name])
                # Aplicar alineación centrada a todas las celdas
                ws_obs.cell(row=fila_actual, column=col_idx).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            
            # Aplicar bordes a cada fila de datos
            for col in range(1, 6):
                celda = ws_obs.cell(row=fila_actual, column=col)
                celda.border = Border(top=thin, bottom=thin, 
                                     left=borde_grueso if col == 1 else thin,
                                     right=borde_grueso if col == 5 else thin)
            
            fila_actual += 1
        
        # Borde inferior grueso para la tabla AM
        for col in range(1, 6):
            celda = ws_obs.cell(row=fila_actual-1, column=col)
            current_border = celda.border
            celda.border = Border(top=current_border.top, bottom=borde_grueso,
                                 left=current_border.left, right=current_border.right)
        
        fila_fin_am = fila_actual - 1
        
        # Colorear tabla AM (usando criterio de 62 dBuV/m)
        for fila in range(fila_inicio_am + 1, fila_fin_am + 1):
            celda_promedio = ws_obs.cell(row=fila, column=3)
            if celda_promedio.value and isinstance(celda_promedio.value, (int, float)):
                valor = float(celda_promedio.value)
                if valor < 62:
                    celda_promedio.fill = rosa
                else:
                    celda_promedio.fill = verde
    
    # Ajustar anchos automáticamente para columnas sin ancho específico
    for col in range(1, ws_obs.max_column + 1):
        col_letter = get_column_letter(col)
        # Solo ajustar si no se ha establecido un ancho específico
        if ws_obs.column_dimensions[col_letter].width is None:
            max_length = 0
            for cell in ws_obs[col_letter]:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws_obs.column_dimensions[col_letter].width = max_length + 2
    
    # Colorear celdas según valores - APLICAR REGLAS DE COLOR MANUALMENTE
    from openpyxl.styles import PatternFill
    
    rojo = PatternFill(start_color="FFFC4A2C", end_color="FFFC4A2C", fill_type="solid")
    amarillo = PatternFill(start_color="FFFFFE9F", end_color="FFFFFE9F", fill_type="solid")
    verde = PatternFill(start_color="FFCDFECE", end_color="FFCDFECE", fill_type="solid")
    rosa = PatternFill(start_color="FFFD9BCB", end_color="FFFD9BCB", fill_type="solid")
    
    # Colorear tabla FM
    if datos_fm is not None and not datos_fm.empty:
        for fila in range(fila_inicio_fm + 1, fila_fin_fm + 1):
            # Colorear Promedio(dBuV/m) - columna C
            celda_promedio = ws_obs.cell(row=fila, column=3)
            if celda_promedio.value and isinstance(celda_promedio.value, (int, float)):
                valor = float(celda_promedio.value)
                if 0 <= valor <= 30:
                    celda_promedio.fill = rojo
                elif 30 < valor < 54:
                    celda_promedio.fill = amarillo
                elif valor >= 54:
                    celda_promedio.fill = verde
            
            # Colorear Ancho de Banda (KHz) - columna D
            celda_ancho = ws_obs.cell(row=fila, column=4)
            if celda_ancho.value and isinstance(celda_ancho.value, (int, float)):
                valor = float(celda_ancho.value)
                if valor <= 220:
                    celda_ancho.fill = verde
                elif valor > 220:
                    celda_ancho.fill = rojo
                    # Agregar observación si es necesario
                    celda_obs = ws_obs.cell(row=fila, column=6)
                    if not celda_obs.value:
                        celda_obs.value = "Opera con ancho de banda mayor a lo autorizado (medición automática)"
                        celda_obs.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    
    # Colorear tabla TV
    if datos_tv is not None and not datos_tv.empty:
        def obtener_banda(freq):
            if (54 <= freq <= 72) or (76 <= freq <= 88):
                return 'I'
            elif 174 <= freq <= 216:
                return 'III'
            elif (470 <= freq <= 488) or (512 <= freq <= 608):
                return 'IV'
            elif 614 <= freq <= 698:
                return 'V'
            else:
                return None
        
        for fila in range(fila_inicio_tv + 1, fila_fin_tv + 1):
            # Obtener frecuencia para determinar banda
            celda_frecuencia = ws_obs.cell(row=fila, column=2)
            celda_promedio = ws_obs.cell(row=fila, column=3)
            celda_obs = ws_obs.cell(row=fila, column=5)
            
            if (celda_frecuencia.value and isinstance(celda_frecuencia.value, (int, float)) and
                celda_promedio.value and isinstance(celda_promedio.value, (int, float))):
                
                frecuencia = float(celda_frecuencia.value)
                valor = float(celda_promedio.value)
                banda = obtener_banda(frecuencia)
                
                if banda == 'I':
                    if valor < 47:
                        celda_promedio.fill = rosa
                        if not celda_obs.value:
                            celda_obs.value = "Niveles por debajo del borde del área de cobertura principal y secundaria"
                    elif 47 <= valor < 68:
                        celda_promedio.fill = amarillo
                    else:
                        celda_promedio.fill = verde
                elif banda == 'III':
                    if valor < 56:
                        celda_promedio.fill = rosa
                        if not celda_obs.value:
                            celda_obs.value = "Niveles por debajo del borde del área de cobertura principal y secundaria"
                    elif 56 <= valor < 71:
                        celda_promedio.fill = amarillo
                    else:
                        celda_promedio.fill = verde
                elif banda in ['IV', 'V']:
                    if valor < 64:
                        celda_promedio.fill = rosa
                        if not celda_obs.value:
                            celda_obs.value = "Niveles por debajo del borde del área de cobertura principal y secundaria"
                    elif 64 <= valor < 74:
                        celda_promedio.fill = amarillo
                    else:
                        celda_promedio.fill = verde
    
    # Ajustar altura de filas para observaciones con texto
    for fila in range(1, ws_obs.max_row + 1):
        for col in range(1, ws_obs.max_column + 1):
            celda = ws_obs.cell(row=fila, column=col)
            if celda.value and "\n" in str(celda.value):
                num_lineas = str(celda.value).count("\n") + 1
                ws_obs.row_dimensions[fila].height = max(ws_obs.row_dimensions[fila].height or 15, num_lineas * 15)

# ------------------ FUNCIÓN PARA OBTENER EL CÓDIGO SEGÚN LA BASE ------------------
def obtener_codigo_base(base):
    """Obtiene el código correspondiente según el nombre de la base"""
    correspondencia = {
        "zamora": "SCS-L01",
        "loja": "SCS-L02", 
        #"cañar": "SCS-L03",  # ñ normal
        #"cañar": "SCS-L03",  # ñ con tilde combinable (n + ˜)
        "tambo":"SCS-L03",
        "macas": "SCS-L04",
        "machala": "SCC-L04",
        "cuenca": "SCS-L05"
    }
    
    # Normalizar el nombre de la base
    base_normalizada = base.lower().strip()
    
    # Manejar diferentes representaciones de "cañar"
    """if (base_normalizada == "cañar" or 
        base_normalizada == "cañar" or  # ñ con tilde combinable
        base_normalizada == "canar" or   # sin tilde
        base_normalizada == "caÃ±ar"):   # posible encoding issue
        base_normalizada = "cañar" """
    
    return correspondencia.get(base_normalizada, f"SCS-{base.upper()}")

# ------------------ FUNCIÓN PARA OBTENER EL NOMBRE DEL MES EN ESPAÑOL ------------------
def obtener_nombre_mes_es(numero_mes):
    """Convierte el número de mes a nombre en español"""
    meses = [
        "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
        "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
    ]
    return meses[numero_mes - 1] if 1 <= numero_mes <= 12 else ""

def procesar_ocupacion(callback_progreso=None, callback_log=None):
    """
    Función para procesar datos de ocupación (placeholder)
    """
    if callback_log:
        callback_log("Iniciando análisis de ocupación...")
        callback_log("Esta funcionalidad está en desarrollo")
    
    # Simular progreso
    for i in range(101):
        if callback_progreso:
            callback_progreso(i)
        time.sleep(0.05)  # Pequeña pausa para simular trabajo
    
    if callback_log:
        callback_log("Análisis de ocupación completado")
    
    return True




# ------------------ FUNCIÓN PRINCIPAL DE PROCESAMIENTO ------------------

def obtener_base(nombre_archivo):
    """Obtiene el nombre base del archivo y filtra 'global'"""
    base = nombre_archivo.split("_")[0].lower().strip()
    # Filtrar "global" y variantes
    if base in ["global", "global", "generico", "general"]:
        return None
    return base

# ... (código anterior sin cambios)

def procesar_datos(callback_progreso=None, callback_log=None, obtener_ciudades=False):
    """
    Función principal que procesa todos los datos con la nueva lógica de cotejo
    """
    # Inicializar directorios
    inicializar_directorios()
    
    # Inicializar lista de ciudades
    ciudades_encontradas = []
    
    # Emitir progreso inicial
    if callback_progreso:
        callback_progreso(0)
    
    # Cargar configuración actual
    config = cargar_configuracion()
    config = limpiar_configuracion_duplicados(config)
    if callback_log:
        callback_log("Iniciando procesamiento de datos...")
        callback_log(f"Ruta FM: {ruta_fm}")
        callback_log(f"Ruta TV: {ruta_tv}")
        callback_log(f"Ruta AM: {ruta_am}")  # ← NUEVO LOG
        callback_log(f"Ruta salida: {ruta_salida}")
    
    # Emitir progreso después de inicialización
    if callback_progreso:
        callback_progreso(5)
    
    # Obtener listas de archivos FILTRANDO "global"
    try:
        archivos_fm = {}
        for f in os.listdir(ruta_fm):
            if f.endswith(".csv"):
                base = obtener_base(f)
                if base:  # Solo agregar si base no es None (filtra "global")
                    archivos_fm[base] = os.path.join(ruta_fm, f)
        
        archivos_tv = {}
        for f in os.listdir(ruta_tv):
            if f.endswith(".csv"):
                base = obtener_base(f)
                if base:  # Solo agregar si base no es None (filtra "global")
                    archivos_tv[base] = os.path.join(ruta_tv, f)
        
    
        # NUEVO: Archivos AM
        archivos_am = {}
        for f in os.listdir(ruta_am):
            if f.endswith(".csv"):
                base = obtener_base(f)
                if base:  # Solo agregar si base no es None (filtra "global")
                    archivos_am[base] = os.path.join(ruta_am, f)
        
        if callback_log:
            callback_log(f"Encontrados {len(archivos_fm)} archivos FM y {len(archivos_tv)} archivos TV (filtrados)")
        
        # Emitir progreso después de leer archivos
        if callback_progreso:
            callback_progreso(10)
            
    except Exception as e:
        if callback_log:
            callback_log(f"Error al leer archivos: {str(e)}")
        return False

    # ✅ NUEVA FUNCIONALIDAD: Cotejar y actualizar frecuencias
    if callback_log:
        callback_log("🔍 Cotejando frecuencias entre config.json y archivos CSV...")
    
        # En la función procesar_datos, busca esta línea y cámbiala:
    frecuencias_a_procesar = cotejar_y_actualizar_frecuencias(config, archivos_fm, archivos_tv, archivos_am, callback_log)  
    
    # Emitir progreso después del cotejo
    if callback_progreso:
        callback_progreso(15)
    
    # Extraer nombres de emisoras de todos los archivos por ciudad (mantener para referencia)
    emisoras_por_ciudad = config.get("emisoras_por_ciudad", {})

    # Procesar archivos FM
    for i, (base, archivo) in enumerate(archivos_fm.items()):
        if not base or not base.strip():
            continue
        base_normalizada = normalizar_nombre_ciudad(base)
        if not base_normalizada:
            continue
        emisoras_fm = extraer_nombres_emisoras(archivo, "FM")
        if base not in emisoras_por_ciudad:
            emisoras_por_ciudad[base] = {"FM": [], "TV": []}
        # Limpiar duplicados y agregar
        emisoras_existentes = {e["nombre"] for e in emisoras_por_ciudad[base]["FM"]}
        for emisora in emisoras_fm:
            if emisora["nombre"] not in emisoras_existentes:
                emisoras_por_ciudad[base]["FM"].append(emisora)
        
        # Emitir progreso incremental para FM
        if callback_progreso:
            progreso_fm = 15 + (i / max(len(archivos_fm), 1)) * 10  # 15% a 25%
            callback_progreso(int(progreso_fm))

    # Procesar archivos TV
    for i, (base, archivo) in enumerate(archivos_tv.items()):
        if not base or not base.strip():
            continue
        base_normalizada = normalizar_nombre_ciudad(base)
        if not base_normalizada:
            continue
        emisoras_tv = extraer_nombres_emisoras(archivo, "TV")
        if base not in emisoras_por_ciudad:
            emisoras_por_ciudad[base] = {"FM": [], "TV": []}
        # Limpiar duplicados y agregar
        emisoras_existentes = {e["nombre"] for e in emisoras_por_ciudad[base]["TV"]}
        for emisora in emisoras_tv:
            if emisora["nombre"] not in emisoras_existentes:
                emisoras_por_ciudad[base]["TV"].append(emisora)
        
        # Emitir progreso incremental para TV
        if callback_progreso:
            progreso_tv = 25 + (i / max(len(archivos_tv), 1)) * 10  # 25% a 35%
            callback_progreso(int(progreso_tv))

        # NUEVO: Procesar archivos AM
    for i, (base, archivo) in enumerate(archivos_am.items()):
        if not base or not base.strip():
            continue
        base_normalizada = normalizar_nombre_ciudad(base)
        if not base_normalizada:
            continue
        emisoras_am = extraer_nombres_emisoras(archivo, "AM")
        if base not in emisoras_por_ciudad:
            emisoras_por_ciudad[base] = {"FM": [], "TV": [], "AM": []}
        
        # CORRECCIÓN: Asegurar que exista la clave 'AM'
        if "AM" not in emisoras_por_ciudad[base]:
            emisoras_por_ciudad[base]["AM"] = []
        
        # Limpiar duplicados y agregar
        emisoras_existentes = {e["nombre"] for e in emisoras_por_ciudad[base]["AM"]}
        for emisora in emisoras_am:
            if emisora["nombre"] not in emisoras_existentes:
                emisoras_por_ciudad[base]["AM"].append(emisora)
        
        # Emitir progreso incremental para AM
        if callback_progreso:
            progreso_am = 35 + (i / max(len(archivos_am), 1)) * 5  # 35% a 40%
            callback_progreso(int(progreso_am))
        
        # Actualizar configuración
        config["emisoras_por_ciudad"] = emisoras_por_ciudad
    
    # Guardar configuración actualizada
    if guardar_configuracion(config):
        total_fm = sum(len(ciudad["FM"]) for ciudad in emisoras_por_ciudad.values() if "FM" in ciudad)
        total_tv = sum(len(ciudad["TV"]) for ciudad in emisoras_por_ciudad.values() if "TV" in ciudad)
        total_am = sum(len(ciudad["AM"]) for ciudad in emisoras_por_ciudad.values() if "AM" in ciudad)  # ← NUEVO
        if callback_log:
            callback_log(f"Guardadas {total_fm} emisoras FM, {total_tv} emisoras TV y {total_am} emisoras AM por ciudad en config.json")
    
    # Emitir progreso después de guardar configuración
    if callback_progreso:
        callback_progreso(40)
    
    nombres_bases = set(archivos_fm.keys()).union(archivos_tv.keys()).union(archivos_am.keys())  # ← AGREGAR AM
    
    # Filtrar bases vacías o nulas
    nombres_bases = {base for base in nombres_bases if base and base.strip()}
    
    # Normalizar TODAS las bases
    ciudades_encontradas = [normalizar_nombre_ciudad(base) for base in nombres_bases]
    ciudades_encontradas = [ciudad for ciudad in ciudades_encontradas if ciudad]  # Filtrar vacíos
    
    if callback_log:
        callback_log(f"Procesando {len(ciudades_encontradas)} bases de datos")
        if obtener_ciudades:
            callback_log(f"Ciudades encontradas: {', '.join(ciudades_encontradas)}")
    
    # Procesar cada base usando el nombre NORMALIZADO
    total_bases = len(nombres_bases)
    for i, base in enumerate(nombres_bases):
        base_normalizada = normalizar_nombre_ciudad(base)
        if not base_normalizada:
            continue
            
        if callback_log:
            callback_log(f"Procesando base: {base_normalizada} (original: {base})")

        # En el procesamiento de cada base, antes del bucle for tipo, archivos, titulo in [...]:
        try:
            # Emitir progreso al iniciar cada base
            if callback_progreso:
                progreso_base = 40 + (i / max(total_bases, 1)) * 55  # 40% a 95%
                callback_progreso(int(progreso_base))
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Informe"
            fila_actual = 1
            
            # Variables para almacenar datos para la hoja de observaciones
            datos_fm = None
            datos_tv = None
            datos_am = None
            
            # NUEVO: Determinar mes_objetivo de manera global para la base
            mes_objetivo = None
            nombre_mes_es = ""
            
            # Buscar el mes objetivo en cualquier archivo disponible (FM, TV o AM)
            for tipo, archivos in [("FM", archivos_fm), ("TV", archivos_tv), ("AM", archivos_am)]:
                if base in archivos:
                    try:
                        df_temp = pd.read_csv(archivos[base], encoding="unicode_escape")
                        df_temp["Tiempo"] = pd.to_datetime(df_temp["Tiempo"], format="%d/%m/%Y  %H:%M:%S,%f", errors='coerce').dt.date
                        mes_temp = df_temp["Tiempo"].dropna().apply(lambda x: x.month).value_counts()
                        if not mes_temp.empty:
                            mes_objetivo = mes_temp.idxmax()
                            nombre_mes_es = ["enero", "febrero", "marzo", "abril", "mayo", "junio", 
                                            "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"][mes_objetivo - 1]
                            break  # Usar el primer mes objetivo encontrado
                    except Exception as e:
                        if callback_log:
                            callback_log(f"⚠️  Error determinando mes para {base} {tipo}: {str(e)}")
                        continue
            
            # Si no se pudo determinar el mes, usar el mes actual
            if mes_objetivo is None:
                mes_objetivo = datetime.now().month
                nombre_mes_es = ["enero", "febrero", "marzo", "abril", "mayo", "junio", 
                                "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"][mes_objetivo - 1]
                if callback_log:
                    callback_log(f"⚠️  No se pudo determinar el mes objetivo para {base}, usando mes actual: {nombre_mes_es}")
            
            # Ahora procesar cada tipo con el mes_objetivo ya definido
            for tipo, archivos, titulo in [
                ("FM", archivos_fm, "FORMULARIO DE CONTROL MENSUAL DE FM"),
                ("TV", archivos_tv, "FORMULARIO DE CONTROL MENSUAL DE TV"),
                ("AM", archivos_am, "FORMULARIO DE CONTROL MENSUAL DE AM")
            ]:
                if base not in archivos:
                    continue

                ruta_archivo = archivos[base]
                df = pd.read_csv(ruta_archivo, encoding="unicode_escape")
                
                # ✅ Filtrar por frecuencias deseadas
                frecuencias_ciudad = frecuencias_a_procesar.get(base, {}).get(tipo, [])
                df = filtrar_dataframe_por_frecuencias(df, frecuencias_ciudad, tipo, callback_log)
                
                if df.empty:
                    if callback_log:
                        callback_log(f"⚠️  No hay datos para procesar en {base} {tipo}")
                    continue

                # CORREGIDO: Renombrar columnas de manera consistente
                if tipo in ["TV", "AM"]:  # ← AM usa el mismo formato que TV
                    df = df.rename(columns={"Nombre de la estación": "ESTACION"})
                
                # CORREGIDO: Seleccionar columnas según el tipo
                if tipo == "FM":
                    df = df[df.columns[:9]]  # FM tiene más columnas
                else:  # TV y AM tienen el mismo formato
                    df = df[df.columns[:5]]

                # Procesamiento común para todos los tipos
                df["ESTACION"] = df["ESTACION"].astype(str).str.strip()
                df = df[df["ESTACION"].notna() & (df["ESTACION"] != "")]
                df["Tiempo"] = pd.to_datetime(df["Tiempo"], format="%d/%m/%Y  %H:%M:%S,%f", errors='coerce').dt.date
                
                # CORREGIDO: Usar el mes_objetivo ya definido globalmente
                # Filtrar por mes objetivo
                df = df[df["Tiempo"].apply(lambda x: x.month == mes_objetivo)]
                df["DIA"] = pd.to_datetime(df["Tiempo"]).dt.day

                # Resto del código sin cambios...
                df["Frecuencia (MHz)"] = df["Frecuencia (Hz)"] / 1_000_000
                df["Level (dBµV/m)"] = df["Level (dBµV/m)"].astype(str).str.replace(",", ".", regex=False).astype(float)

                # Crear pivot table
                agrupado = df.groupby(["ESTACION", "Frecuencia (MHz)", "DIA"])["Level (dBµV/m)"].mean().reset_index()
                pivot = agrupado.pivot(index=["ESTACION", "Frecuencia (MHz)"], columns="DIA", values="Level (dBµV/m)")
                todos_los_dias = list(range(1, 32))
                pivot = pivot.reindex(columns=todos_los_dias, fill_value=0)

                pivot[todos_los_dias] = pivot[todos_los_dias].replace(0, "-")
                pivot_numeric = pivot[todos_los_dias].replace("-", pd.NA).apply(pd.to_numeric, errors="coerce")
                pivot["Promedio(dBuV/m)"] = pivot_numeric.mean(axis=1, skipna=True).round(2)

                # Resto del código sin cambios...

                # CORREGIDO: Procesamiento específico por tipo
                if tipo == "FM":
                    ancho_banda = (
                        df.groupby(["ESTACION", "Frecuencia (MHz)"])["Bandwidth (Hz)"].mean().reset_index()
                    )
                    ancho_banda["Ancho de Banda (KHz)"] = (ancho_banda["Bandwidth (Hz)"] / 1000).round(2)
                    pivot = pivot.reset_index().merge(ancho_banda.drop(columns=["Bandwidth (Hz)"]), on=["ESTACION", "Frecuencia (MHz)"], how="left")
                else:  # TV y AM
                    pivot = pivot.reset_index()

                # Agregar columnas manuales y observaciones
                pivot["Medición Manual"] = ""
                pivot["OBSERVACIONES"] = ""
                
                # CORREGIDO: Guardar datos para hoja de observaciones de manera separada
                if tipo == "FM":
                    datos_fm = pivot[["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", "Ancho de Banda (KHz)", "Medición Manual", "OBSERVACIONES"]].copy()
                elif tipo == "TV":
                    datos_tv = pivot[["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", "Medición Manual", "OBSERVACIONES"]].copy()
                elif tipo == "AM":
                    datos_am = pivot[["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", "Medición Manual", "OBSERVACIONES"]].copy()

                # Ordenar y preparar datos para exportación
                pivot = pivot.sort_values(by="Frecuencia (MHz)")
                pivot = pivot.reset_index(drop=True)

                # Redondear valores numéricos
                for col in pivot.select_dtypes(include="number").columns:
                    pivot[col] = pivot[col].round(2)

                # Escribir datos en la hoja
                for i, row in enumerate(dataframe_to_rows(pivot, index=False, header=True)):
                    for j, val in enumerate(row, start=1):
                        ws.cell(row=fila_actual + i, column=j, value=val)

                # CORREGIDO: Encabezados específicos por tipo
                if tipo == "FM":
                    encabezado = [
                        "INFORME DE CONTROL TÉCNICO",
                        "No. IT-CZ06-R-2025-00XX",
                        "AGENCIA DE REGULACIÓN Y CONTROL DE LAS TELECOMUNICACIONES",
                        "COORDINACIÓN ZONAL 6",
                        "ESTACIÓN DE COMPROBACIÓN TÉCNICA",
                        titulo,
                        "CIUDAD:" + base.upper(),
                        f"PERIODO: {nombre_mes_es.upper()}",
                        f"FECHA PRESENTACIÓN: {fecha_actual}"
                    ]
                    insertar_imagenes(ws, fila_actual + 4, tipo="FM")
                else:  # TV y AM
                    encabezado = [
                        "AGENCIA DE REGULACIÓN Y CONTROL DE LAS TELECOMUNICACIONES",
                        "COORDINACIÓN ZONAL 6",
                        "ESTACIÓN DE COMPROBACIÓN TÉCNICA",
                        titulo,
                        "CIUDAD:" + base.upper(),
                        f"PERIODO: {nombre_mes_es.upper()}",
                        f"FECHA PRESENTACIÓN: {fecha_actual}"
                    ]
                    insertar_imagenes(ws, fila_actual + 2, tipo=tipo)  # ← Usar el tipo actual

                # Aplicar formato
                formatear_hoja(ws, fila_actual, encabezado)

                # Colorear celdas
                fila_encabezados_columnas = fila_actual + len(encabezado)
                fila_inicio_datos = fila_encabezados_columnas + 1
                fila_fin_datos = fila_inicio_datos + len(pivot) - 1

                # CORREGIDO: Columnas a colorear según el tipo
                if tipo == "FM":
                    columnas_colorear = ["Promedio (dBuV/m)", "Ancho de Banda\n(KHz)", *list(range(1, 32))]
                else:  # TV y AM
                    columnas_colorear = ["Promedio (dBuV/m)", *list(range(1, 32))]
                
                colorear_celdas_por_valor(ws, fila_inicio_datos, fila_fin_datos, tipo, columnas_colorear)

                fila_actual = ws.max_row + 3

            # Crear hojas adicionales
            wb.create_sheet("Manual FM")
            wb.create_sheet("Manual TV")
            crear_hoja_observaciones(wb, datos_fm, datos_tv, datos_am)
            
            # Reordenar hojas
            orden_hojas = ["Informe", "Manual FM", "Manual TV", "Observaciones"]
            for hoja in orden_hojas:
                if hoja in wb.sheetnames:
                    wb.move_sheet(hoja, -len(orden_hojas))
                    orden_hojas.remove(hoja)
            
            codigo_base = obtener_codigo_base(base)
            nombre_ciudad = base.upper()
            nombre_mes_completo = obtener_nombre_mes_es(mes_objetivo)
            nombre_salida = f"{codigo_base}_Procesamiento{nombre_ciudad}_{nombre_mes_completo}2025.xlsx"
            
            # Combinar observaciones para TV
            for sheet in wb.worksheets:
                if "Informe" not in sheet.title:
                    continue

                # Buscar fila del segundo encabezado (TV)
                for fila in range(1, sheet.max_row + 1):
                    val = sheet.cell(row=fila, column=1).value
                    if isinstance(val, str) and "FORMULARIO DE CONTROL MENSUAL DE TV" in val:
                        fila_encabezado_tv = fila
                        break
                else:
                    continue  # No hay TV en esta hoja

                # Buscar fila exacta de la cabecera de la tabla (donde dice "ESTACION")
                fila_tabla_tv = None
                for fila in range(fila_encabezado_tv + 1, sheet.max_row + 1):
                    if sheet.cell(row=fila, column=1).value == "ESTACION":
                        fila_tabla_tv = fila
                        break

                if fila_tabla_tv is None:
                    continue  # Seguridad por si no se encuentra

                ultima_fila_tv = sheet.max_row
                col_final = sheet.max_column
                col_penultima = col_final - 1

                # Obtener anchos actuales
                ancho_col1 = sheet.column_dimensions[get_column_letter(col_penultima)].width
                ancho_col2 = sheet.column_dimensions[get_column_letter(col_final)].width
                ancho_combinado = (ancho_col1 or 10) + (ancho_col2 or 10)

                # Combinar celdas del cuerpo de la tabla
                for fila in range(fila_tabla_tv + 1, ultima_fila_tv + 1):
                    sheet.merge_cells(start_row=fila, start_column=col_penultima, end_row=fila, end_column=col_final)

                # Combinar encabezado de tabla ("OBSERVACIONES")
                sheet.merge_cells(start_row=fila_tabla_tv, start_column=col_penultima, end_row=fila_tabla_tv, end_column=col_final)
                celda_obs = sheet.cell(row=fila_tabla_tv, column=col_penultima)
                celda_obs.value = "OBSERVACIONES"
                celda_obs.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                # Asignar nuevo ancho
                sheet.column_dimensions[get_column_letter(col_penultima)].width = ancho_combinado
                ws.sheet_view.showGridLines = False

            wb.save(os.path.join(ruta_salida, nombre_salida))
            
            if callback_log:
                callback_log(f"✅ Archivo generado: {nombre_salida}")
                
        except Exception as e:
            if callback_log:
                callback_log(f"❌ Error procesando base {base}: {str(e)}")
    
    # Emitir progreso final
    if callback_progreso:
        callback_progreso(100)
        
    if callback_log:
        callback_log("Procesamiento completado")
    
    # Devolver resultado y lista de ciudades si se solicitó
    if obtener_ciudades:
        return True, ciudades_encontradas
    else:
        return True

# ------------------ EJECUCIÓN DIRECTA (para testing) ------------------

if __name__ == "__main__":
    # Si se ejecuta directamente, usar callbacks simples
    def mostrar_progreso(progreso):
        print(f"Progreso: {progreso}%")
    
    def mostrar_log(mensaje):
        print(mensaje)
    
    # Procesar datos
    resultado, ciudades = procesar_datos(mostrar_progreso, mostrar_log, obtener_ciudades=True)
    
    if resultado:
        print("Procesamiento completado con éxito")
        print(f"Ciudades encontradas: {ciudades}")
    else:
        print("Ocurrieron errores durante el procesamiento")
