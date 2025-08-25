import pandas as pd
import calendar
import os

# -------- CONFIGURACIÓN --------
ruta_entrada = "MedicionesCSV"
archivo_entrada = "canar_juliototal.csv"  # ⚠️ Cambia el nombre si usas otro archivo
ruta_salida = "pruebas"
os.makedirs(ruta_salida, exist_ok=True)

# Construir rutas
ruta_csv_entrada = os.path.join(ruta_entrada, archivo_entrada)

# Cargar datos
df = pd.read_csv(ruta_csv_entrada, encoding='unicode_escape')
df = df.iloc[:, :9]
df["ESTACION"] = df["ESTACION"].astype(str).str.strip()
df = df[df["ESTACION"].notna() & (df["ESTACION"] != "")]

# Formatear columnas
df["Tiempo"] = pd.to_datetime(df["Tiempo"], format="%d/%m/%Y  %H:%M:%S,%f", errors='coerce').dt.date

# Detectar mes con más datos
mes_ocurrencias = df["Tiempo"].apply(lambda x: x.month)
mes_objetivo = mes_ocurrencias.value_counts().idxmax()

nombre_mes_es = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril",
    5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
    9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"
}[mes_objetivo]

# Filtrar por mes
df["Frecuencia (MHz)"] = df["Frecuencia (Hz)"] / 1_000_000
df["Level (dBµV/m)"] = df["Level (dBµV/m)"].astype(str).str.replace(",", ".", regex=False).astype(float)
df = df[df["Tiempo"].apply(lambda x: x.month == mes_objetivo)]

if df.empty:
    print(f"⚠ No hay datos válidos para el mes más frecuente ({nombre_mes_es}) en {archivo_entrada}")
else:
    df["DIA"] = pd.to_datetime(df["Tiempo"]).dt.day

    agrupado = df.groupby(["ESTACION", "Frecuencia (MHz)", "DIA"])["Level (dBµV/m)"].mean().reset_index()
    pivot = agrupado.pivot(index=["ESTACION", "Frecuencia (MHz)"], columns="DIA", values="Level (dBµV/m)")

    # Rellenar días del 1 al 31
    todos_los_dias = list(range(1, 32))
    pivot = pivot.reindex(columns=todos_los_dias, fill_value=0)

    # Promedio mensual Level
    promedios_sin_ceros = pivot.replace(0, pd.NA).mean(axis=1, skipna=True)
    pivot["Promedio Mensual"] = promedios_sin_ceros.astype(float).round(3)

    # Promedio mensual Bandwidth en kHz
    bandwidth_promedios = (
        df.groupby(["ESTACION", "Frecuencia (MHz)"])["Bandwidth (Hz)"]
        .mean()
        .reset_index()
    )
    bandwidth_promedios["Promedio Bandwidth (kHz)"] = (bandwidth_promedios["Bandwidth (Hz)"] / 1000).round(3)
    bandwidth_promedios = bandwidth_promedios.drop(columns=["Bandwidth (Hz)"])
    pivot = pivot.merge(bandwidth_promedios, on=["ESTACION", "Frecuencia (MHz)"], how="left")

    # Ordenar
    pivot = pivot.sort_values(by="Frecuencia (MHz)")

    # Reemplazar ceros con guiones
    pivot[todos_los_dias] = pivot[todos_los_dias].replace(0, "-")

    # Columnas manuales vacías
    pivot["Medicion Manual"] = ""
    pivot["Observaciones"] = ""

    # Resetear índice y eliminar columna 'index' si aparece
    pivot = pivot.reset_index(drop=True)

    # Generar nombre del archivo de salida
    base_nombre = archivo_entrada.split("_")[0]
    nombre_excel = f"{base_nombre}_promediostotal.xlsx"
    ruta_excel_salida = os.path.join(ruta_salida, nombre_excel)

    # Redondear columnas numéricas
    for col in pivot.select_dtypes(include="number").columns:
         pivot[col] = pivot[col].round(3)

    with pd.ExcelWriter(ruta_excel_salida, engine="openpyxl") as writer:
        pivot.to_excel(writer, index=False, sheet_name="Resumen")

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
import os


# Abrir el archivo
wb = load_workbook(ruta_excel_salida)
ws = wb["Resumen"] # o  si sabes el nombre de la hoja

# Encabezado
encabezado_lineas = [
    "INFORME DE CONTROL TÉCNICO",
    "No. IT-CZ06-R-2025-00XX",
    "AGENCIA DE REGULACIÓN Y CONTROL DE LAS TELECOMUNICACIONES",
    "COORDINACIÓN ZONAL 6",
    "ESTACIÓN DE COMPROBACIÓN TÉCNICA",
    "FORMULARIO DE CONTROL MENSUAL DE FM",
    "CIUDAD: [CIUDAD]",
    "PERIODO: JULIO DE 2025",
    "FECHA PRESENTACIÓN: 06/08/2025"
]

num_columnas = ws.max_column
ultima_col = get_column_letter(num_columnas)

# Insertar encabezado
ws.insert_rows(1, amount=len(encabezado_lineas))
for i, texto in enumerate(encabezado_lineas, start=1):
    celda = ws.cell(row=i, column=1, value=texto)
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=num_columnas)
    celda.alignment = Alignment(horizontal="center", vertical="center")
    celda.font = Font(bold=True, size=12)

# Quitar bordes del encabezado
for row in ws.iter_rows(min_row=1, max_row=len(encabezado_lineas), min_col=1, max_col=num_columnas):
    for cell in row:
        cell.border = Border()

# Insertar imágenes
ruta_imagenes = "Img"
logo_izquierda = os.path.join(ruta_imagenes, "ARCOTEL.png")
logo_derecha = os.path.join(ruta_imagenes, "nEcuador.png")

if os.path.exists(logo_izquierda):
    img_left = XLImage(logo_izquierda)
    img_left.width = 120
    img_left.height = 60
    img_left.left = 1000
    img_left.top = 500
    ws.add_image(img_left, "A1")

if os.path.exists(logo_derecha):
    img_right = XLImage(logo_derecha)
    img_right.width = 120
    img_right.height = 60
    img_right.left = -1000
    img_right.top = 500
    ws.add_image(img_right, f"{ultima_col}1")

# Ajustar ancho de columnas
for i in range(1, num_columnas + 1):
    col_letter = get_column_letter(i)
    max_length = 0
    for cell in ws[col_letter]:
        try:
            max_length = max(max_length, len(str(cell.value)))
        except:
            pass
    ws.column_dimensions[col_letter].width = max_length + 2

# Agregar marco exterior con líneas más gruesas
borde_grueso = Side(border_style="medium")
inicio_fila_tabla = len(encabezado_lineas) + 1
fin_fila_tabla = ws.max_row
for row in range(inicio_fila_tabla, fin_fila_tabla + 1):
    for col in range(1, num_columnas + 1):
        cell = ws.cell(row=row, column=col)
        borde = Border()
        if row == inicio_fila_tabla:
            borde += Border(top=borde_grueso)
        if row == fin_fila_tabla:
            borde += Border(bottom=borde_grueso)
        if col == 1:
            borde += Border(left=borde_grueso)
        if col == num_columnas:
            borde += Border(right=borde_grueso)
        cell.border = cell.border + borde

# Guardar
wb.save(ruta_excel_salida)
print("✅ Encabezado + imágenes + ancho ajustado correctamente.")

print(f"✅ Archivo generado en formato Excel: {ruta_excel_salida}")
