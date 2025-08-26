import pandas as pd
import calendar
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

# -------- CONFIGURACIÓN --------
archivo_entrada = "TvH025_julio.csv"
ruta_salida = "pruebas"
ruta_imagenes = "Img"
os.makedirs(ruta_salida, exist_ok=True)
fecha_actual = datetime.today().strftime("%d/%m/%Y")

# Leer archivo y dejar solo 5 columnas
df = pd.read_csv(archivo_entrada, encoding='unicode_escape')
df = df.iloc[:, :5]

# Renombrar campo estación
df = df.rename(columns={"Nombre de la estación": "ESTACION"})

# Limpiar estación
df["ESTACION"] = df["ESTACION"].astype(str).str.strip()
df = df[df["ESTACION"].notna() & (df["ESTACION"] != "")]

# Formato de fecha
df["Tiempo"] = pd.to_datetime(df["Tiempo"], format="%d/%m/%Y  %H:%M:%S,%f", errors='coerce').dt.date

# Detectar mes con más datos
mes_ocurrencias = df["Tiempo"].apply(lambda x: x.month)
mes_objetivo = mes_ocurrencias.value_counts().idxmax()
nombre_mes_es = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio",
    7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"
}[mes_objetivo]

# Procesamiento de datos
df["Frecuencia (MHz)"] = df["Frecuencia (Hz)"] / 1_000_000
df["Level (dBµV/m)"] = df["Level (dBµV/m)"].astype(str).str.replace(",", ".", regex=False).astype(float)
df = df[df["Tiempo"].apply(lambda x: x.month == mes_objetivo)]
df["DIA"] = pd.to_datetime(df["Tiempo"]).dt.day

# Agrupación y pivot
agrupado = df.groupby(["ESTACION", "Frecuencia (MHz)", "DIA"])["Level (dBµV/m)"].mean().reset_index()
pivot = agrupado.pivot(index=["ESTACION", "Frecuencia (MHz)"], columns="DIA", values="Level (dBµV/m)")
todos_los_dias = list(range(1, 32))
pivot = pivot.reindex(columns=todos_los_dias, fill_value=0)
pivot["Promedio(dBuV/m)"] = pivot.replace(0, pd.NA).mean(axis=1, skipna=True).astype(float).round(2)
pivot["Medición Manual"] = ""
pivot["OBSERVACIONES"] = ""
pivot = pivot.sort_values(by="Frecuencia (MHz)").reset_index()

# Guardar a Excel
nombre_excel = "TvH025_promediostotal.xlsx"
ruta_excel_salida = os.path.join(ruta_salida, nombre_excel)
for col in pivot.select_dtypes(include="number").columns:
    pivot[col] = pivot[col].round(2)
with pd.ExcelWriter(ruta_excel_salida, engine="openpyxl") as writer:
    pivot.to_excel(writer, index=False, sheet_name="Resumen")

# ---------------- FORMATO EXCEL ----------------
wb = load_workbook(ruta_excel_salida)
ws = wb["Resumen"]
ws.sheet_view.showGridLines = False

# Encabezado
encabezado_lineas = [
    "INFORME DE CONTROL TÉCNICO",
    "No. IT-CZ06-R-2025-00XX",
    "AGENCIA DE REGULACIÓN Y CONTROL DE LAS TELECOMUNICACIONES",
    "COORDINACIÓN ZONAL 6",
    "ESTACIÓN DE COMPROBACIÓN TÉCNICA",
    "FORMULARIO DE CONTROL MENSUAL DE FM",
    "CIUDAD: TVH025",
    f"PERIODO: {nombre_mes_es.upper()}",
    f"FECHA PRESENTACIÓN: {fecha_actual}"
]
num_columnas = ws.max_column
ws.insert_rows(1, amount=len(encabezado_lineas))
for i, texto in enumerate(encabezado_lineas, start=1):
    celda = ws.cell(row=i, column=1, value=texto)
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=num_columnas)
    celda.alignment = Alignment(horizontal="center", vertical="center")
    celda.font = Font(bold=True, size=12)

# Insertar imágenes
logo_izquierda = os.path.join(ruta_imagenes, "ARCOTEL.png")
logo_derecha = os.path.join(ruta_imagenes, "nEcuador.png")
if os.path.exists(logo_izquierda):
    img_left = XLImage(logo_izquierda)
    img_left.width = 500
    img_left.height = 100
    img_left.left = 10000
    img_left.top = 500
    ws.add_image(img_left, "A4")
if os.path.exists(logo_derecha):
    img_right = XLImage(logo_derecha)
    img_right.width = 260
    img_right.height = 120
    img_right.left = 1000
    img_right.top = 500
    ws.add_image(img_right, "AH3")

# Ancho columnas
for i in range(1, num_columnas + 1):
    col_letter = get_column_letter(i)
    max_length = 0
    for cell in ws[col_letter]:
        try:
            max_length = max(max_length, len(str(cell.value)))
        except:
            pass
    ws.column_dimensions[col_letter].width = max_length + 2

# Bordes
inicio_fila_tabla = len(encabezado_lineas) + 1
fin_fila_tabla = ws.max_row
thin = Side(border_style="thin")
for row in range(inicio_fila_tabla, fin_fila_tabla + 1):
    for col in range(1, num_columnas + 1):
        ws.cell(row=row, column=col).border = Border(top=thin, bottom=thin, left=thin, right=thin)
borde_grueso = Side(border_style="medium")
for row in range(inicio_fila_tabla, fin_fila_tabla + 1):
    for col in range(1, num_columnas + 1):
        cell = ws.cell(row=row, column=col)
        current = cell.border
        cell.border = Border(
            top=borde_grueso if row == inicio_fila_tabla else current.top,
            bottom=borde_grueso if row == fin_fila_tabla else current.bottom,
            left=borde_grueso if col == 1 else current.left,
            right=borde_grueso if col == num_columnas else current.right
        )

# Alineación centrada
alineacion = Alignment(horizontal="center", vertical="center")
for row in range(inicio_fila_tabla, fin_fila_tabla + 1):
    for col in range(1, num_columnas + 1):
        ws.cell(row=row, column=col).alignment = alineacion

# Saltos de línea para cabeceras específicas
for col in range(1, num_columnas + 1):
    cell = ws.cell(row=inicio_fila_tabla, column=col)
    if cell.value == "Promedio(dBuV/m)":
        cell.value = "Promedio\n(dBuV/m)"
    elif cell.value == "Medición Manual":
        cell.value = "Medición Manual\nAB(KHz) o NIVEL\n(dBµV/m)"
    cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
ws.row_dimensions[inicio_fila_tabla].height = 45

# Guardar archivo
wb.save(ruta_excel_salida)
print(f"✅ Archivo generado en formato Excel: {ruta_excel_salida}")
