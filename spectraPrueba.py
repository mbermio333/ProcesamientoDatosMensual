import openpyxl
import json
from datetime import datetime

def filtrar_datos_arcotel(archivo_excel, archivo_json):
    """
    Filtra datos de la base de datos ARCOTEL según criterios específicos y guarda en JSON
    """
    try:
        # Cargar el archivo Excel
        workbook = openpyxl.load_workbook(archivo_excel)
        sheet = workbook['descarga']  # Usar la hoja 'descarga'
        
        # Definir los filtros
        filtros = {
            'PROVINCIA_A': 'ZAMORA CHINCHIPE',
            'SERVICIO': 'FM - Frecuencia Modulada', 
            'AREAS_OP_MR': 'ZAMORA'
        }
        
        # Mapeo de columnas a extraer
        columnas_extraer = {
            'NOMBRES': 'C',
            'RED': 'E', 
            'FRECUENCIA': 'F',
            'ESTADO_ESTACION': 'AI',
            'ESTACION': 'AR',
            'ESTADO_SOLICITUD': 'BC',
            'SUSCRIPCION': 'AJ',
            'VIGENCIA': 'AK'
        }
        
        # Obtener índices de las columnas de filtro
        col_provincia = 'A'  # PROVINCIA_A
        col_servicio = 'B'   # SERVICIO
        col_areas_op = 'I'   # AREAS_OP_MR
        
        # Lista para almacenar los resultados
        resultados = []
        
        # Iterar sobre las filas (empezando desde la fila 2 para saltar los encabezados)
        for row in range(2, sheet.max_row + 1):
            # Obtener valores de las columnas de filtro
            provincia = sheet[f'{col_provincia}{row}'].value
            servicio = sheet[f'{col_servicio}{row}'].value
            areas_op = sheet[f'{col_areas_op}{row}'].value
            
            # Aplicar filtros
            if (provincia == filtros['PROVINCIA_A'] and 
                servicio == filtros['SERVICIO'] and 
                areas_op and filtros['AREAS_OP_MR'] in str(areas_op)):
                
                # Crear diccionario con los datos a extraer
                registro = {}
                
                for nombre_col, col in columnas_extraer.items():
                    valor = sheet[f'{col}{row}'].value
                    
                    # Convertir fechas a string si es necesario
                    if isinstance(valor, datetime):
                        valor = valor.strftime('%Y-%m-%d %H:%M:%S')
                    elif valor is None:
                        valor = ""
                    
                    registro[nombre_col] = valor
                
                resultados.append(registro)
        
        # Guardar resultados en archivo JSON
        with open(archivo_json, 'w', encoding='utf-8') as json_file:
            json.dump(resultados, json_file, ensure_ascii=False, indent=2)
        
        print(f"Proceso completado. Se encontraron {len(resultados)} registros.")
        print(f"Resultados guardados en: {archivo_json}")
        
        # Mostrar preview de los resultados
        if resultados:
            print("\nPreview de los primeros registros:")
            for i, registro in enumerate(resultados[:3], 1):
                print(f"Registro {i}: {registro}")
        
        return resultados
        
    except FileNotFoundError:
        print(f"Error: No se encontró el archivo {archivo_excel}")
        return []
    except Exception as e:
        print(f"Error durante el procesamiento: {str(e)}")
        return []

# Función alternativa si necesitas más flexibilidad en los filtros
def filtrar_datos_flexible(archivo_excel, archivo_json, filtros_personalizados=None):
    """
    Versión más flexible que permite personalizar los filtros
    """
    try:
        workbook = openpyxl.load_workbook(archivo_excel)
        sheet = workbook['descarga']
        
        # Usar filtros por defecto o los personalizados
        if filtros_personalizados is None:
            filtros = {
                'PROVINCIA_A': 'ZAMORA CHINCHIPE',
                'SERVICIO': 'FM - Frecuencia Modulada',
                'AREAS_OP_MR': 'ZAMORA'
            }
        else:
            filtros = filtros_personalizados
        
        # Mapeo de columnas
        columnas_extraer = {
            'NOMBRES': 'C',
            'RED': 'E', 
            'FRECUENCIA': 'F',
            'ESTADO_ESTACION': 'AI',
            'ESTACION': 'AR',
            'ESTADO_SOLICITUD': 'BC',
            'SUSCRIPCION': 'AJ',
            'VIGENCIA': 'AK'
        }
        
        resultados = []
        
        for row in range(2, sheet.max_row + 1):
            cumple_filtros = True
            
            # Verificar cada filtro
            if 'PROVINCIA_A' in filtros:
                provincia = sheet[f'A{row}'].value
                if provincia != filtros['PROVINCIA_A']:
                    cumple_filtros = False
            
            if 'SERVICIO' in filtros and cumple_filtros:
                servicio = sheet[f'B{row}'].value
                if servicio != filtros['SERVICIO']:
                    cumple_filtros = False
            
            if 'AREAS_OP_MR' in filtros and cumple_filtros:
                areas_op = sheet[f'I{row}'].value
                if not areas_op or filtros['AREAS_OP_MR'] not in str(areas_op):
                    cumple_filtros = False
            
            if cumple_filtros:
                registro = {}
                for nombre_col, col in columnas_extraer.items():
                    valor = sheet[f'{col}{row}'].value
                    
                    if isinstance(valor, datetime):
                        valor = valor.strftime('%Y-%m-%d %H:%M:%S')
                    elif valor is None:
                        valor = ""
                    
                    registro[nombre_col] = valor
                
                resultados.append(registro)
        
        # Guardar en JSON
        with open(archivo_json, 'w', encoding='utf-8') as json_file:
            json.dump(resultados, json_file, ensure_ascii=False, indent=2)
        
        print(f"Proceso completado. Se encontraron {len(resultados)} registros.")
        return resultados
        
    except Exception as e:
        print(f"Error: {str(e)}")
        return []

if __name__ == "__main__":
    # Configuración
    archivo_excel = "9. SPECTRA_10sep2025.xlsx"
    archivo_json = "resultados_filtrados.json"
    
    # Ejecutar el filtrado
    resultados = filtrar_datos_arcotel(archivo_excel, archivo_json)
    
    # Si quieres usar filtros personalizados:
    # filtros_personalizados = {
    #     'PROVINCIA_A': 'EL ORO',
    #     'SERVICIO': 'AM - Amplitud Modulada',
    #     'AREAS_OP_MR': 'MACHALA'
    # }
    # resultados = filtrar_datos_flexible(archivo_excel, "resultados_personalizados.json", filtros_personalizados)