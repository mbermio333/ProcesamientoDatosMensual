# -*- coding: utf-8 -*-
import sys
import os
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image as XLImage

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, 
    QHBoxLayout, QPushButton, QTableWidget, QTableWidgetItem,
    QTextEdit, QFileDialog, QMessageBox, QLabel, QTabWidget,
    QStatusBar, QAction, QToolBar, QGroupBox, QLineEdit,
    QProgressBar, QCheckBox
)
from PyQt5.QtCore import Qt
from PyQt5.QtGui import QIcon

class DatabaseApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.initUI()
        
    def initUI(self):
        # Configuración básica de la ventana
        self.setWindowTitle("Procesador de Base de Datos - ARCOTEL")
        self.setGeometry(100, 100, 1000, 700)
        
        # Widget central
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        
        # Layout principal
        main_layout = QVBoxLayout(central_widget)
        
        # Crear pestañas
        self.tabs = QTabWidget()
        main_layout.addWidget(self.tabs)
        
        # Pestaña 1: Configuración
        self.setup_config_tab()
        
        # Pestaña 2: Procesamiento
        self.setup_process_tab()
        
        # Barra de estado
        self.status_bar = QStatusBar()
        self.setStatusBar(self.status_bar)
        self.status_bar.showMessage("Listo para comenzar")
        
        # Mostrar ventana
        self.show()
    
    def setup_config_tab(self):
        # Crear widget para la pestaña de configuración
        config_tab = QWidget()
        layout = QVBoxLayout(config_tab)
        
        # Grupo para rutas
        paths_group = QGroupBox("Configuración de Rutas")
        paths_layout = QVBoxLayout()
        
        # Ruta FM
        fm_layout = QHBoxLayout()
        fm_layout.addWidget(QLabel("Ruta Mediciones FM:"))
        self.fm_path = QLineEdit("MedicionesFmCSV")
        fm_layout.addWidget(self.fm_path)
        self.btn_fm_browse = QPushButton("Examinar")
        self.btn_fm_browse.clicked.connect(lambda: self.browse_folder(self.fm_path))
        fm_layout.addWidget(self.btn_fm_browse)
        paths_layout.addLayout(fm_layout)
        
        # Ruta TV
        tv_layout = QHBoxLayout()
        tv_layout.addWidget(QLabel("Ruta Mediciones TV:"))
        self.tv_path = QLineEdit("MedicionesTvCSV")
        tv_layout.addWidget(self.tv_path)
        self.btn_tv_browse = QPushButton("Examinar")
        self.btn_tv_browse.clicked.connect(lambda: self.browse_folder(self.tv_path))
        tv_layout.addWidget(self.btn_tv_browse)
        paths_layout.addLayout(tv_layout)
        
        # Ruta Salida
        output_layout = QHBoxLayout()
        output_layout.addWidget(QLabel("Ruta de Salida:"))
        self.output_path = QLineEdit("ReportesUnificados")
        output_layout.addWidget(self.output_path)
        self.btn_output_browse = QPushButton("Examinar")
        self.btn_output_browse.clicked.connect(lambda: self.browse_folder(self.output_path))
        output_layout.addWidget(self.btn_output_browse)
        paths_layout.addLayout(output_layout)
        
        # Ruta Imágenes
        img_layout = QHBoxLayout()
        img_layout.addWidget(QLabel("Ruta de Imágenes:"))
        self.img_path = QLineEdit("Img")
        img_layout.addWidget(self.img_path)
        self.btn_img_browse = QPushButton("Examinar")
        self.btn_img_browse.clicked.connect(lambda: self.browse_folder(self.img_path))
        img_layout.addWidget(self.btn_img_browse)
        paths_layout.addLayout(img_layout)
        
        paths_group.setLayout(paths_layout)
        layout.addWidget(paths_group)
        
        # Botón de procesamiento
        self.btn_start = QPushButton("Iniciar Procesamiento")
        self.btn_start.clicked.connect(self.start_processing)
        layout.addWidget(self.btn_start)
        
        # Añadir pestaña al tab widget
        self.tabs.addTab(config_tab, "Configuración")
    
    def setup_process_tab(self):
        # Crear widget para la pestaña de procesamiento
        process_tab = QWidget()
        layout = QVBoxLayout(process_tab)
        
        # Barra de progreso
        self.progress_bar = QProgressBar()
        self.progress_bar.setValue(0)
        layout.addWidget(self.progress_bar)
        
        # Área de log
        layout.addWidget(QLabel("Log de Procesamiento:"))
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        layout.addWidget(self.log_text)
        
        # Añadir pestaña al tab widget
        self.tabs.addTab(process_tab, "Procesamiento")
    
    def browse_folder(self, line_edit):
        folder = QFileDialog.getExistingDirectory(self, "Seleccionar Carpeta")
        if folder:
            line_edit.setText(folder)
    
    def start_processing(self):
        self.log_text.append("Iniciando procesamiento...")
        self.status_bar.showMessage("Procesando...")
        self.process_files()
    
    # ------------------ FUNCIONES ORIGINALES ------------------
    
    def combinar_observaciones_solo_en_tv(self, ws, fila_inicio_tabla, fila_fin_tabla, num_columnas, es_tv=True):
        """Combina horizontalmente la última y penúltima columna para cada fila SOLO si es TV."""
        if es_tv:
            for fila in range(fila_inicio_tabla + 1, fila_fin_tabla + 1):
                ws.merge_cells(start_row=fila, start_column=num_columnas - 1, end_row=fila, end_column=num_columnas)
            
            ws.merge_cells(start_row=fila_inicio_tabla, start_column=num_columnas - 1, end_row=fila_inicio_tabla, end_column=num_columnas)
            celda = ws.cell(row=fila_inicio_tabla, column=num_columnas - 1)
            celda.value = "OBSERVACIONES"
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.cell(row=fila_inicio_tabla, column=num_columnas).value = None

    def formatear_hoja(self, ws, fila_inicio, encabezado_lineas):
        num_columnas = ws.max_column
        ws.insert_rows(fila_inicio, amount=len(encabezado_lineas))

        for i, texto in enumerate(encabezado_lineas):
            fila = fila_inicio + i
            celda = ws.cell(row=fila, column=1, value=texto)
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=num_columnas)
            celda.alignment = Alignment(horizontal="center", vertical="center")
            celda.font = Font(bold=True, size=12)

        # Insertar imágenes
        logo_izquierda = os.path.join(self.img_path.text(), "ARCOTEL.png")
        logo_derecha = os.path.join(self.img_path.text(), "nEcuador.png")

        fila_img_izq = fila_inicio + 2
        fila_img_der = fila_inicio + 2

        if os.path.exists(logo_izquierda):
            img_left = XLImage(logo_izquierda)
            img_left.width = 500
            img_left.height = 100
            ws.add_image(img_left, f"A{fila_img_izq}")

        if os.path.exists(logo_derecha):
            img_right = XLImage(logo_derecha)
            img_right.width = 260
            img_right.height = 120
            ws.add_image(img_right, f"AH{fila_img_der}")

        # Bordes
        inicio_fila_tabla = fila_inicio + len(encabezado_lineas)
        fin_fila_tabla = ws.max_row
        thin = Side(border_style="thin")
        borde_grueso = Side(border_style="medium")

        for row in range(inicio_fila_tabla, fin_fila_tabla + 1):
            for col in range(1, num_columnas + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                if row == inicio_fila_tabla:
                    cell.border = Border(top=borde_grueso, bottom=thin,
                                         left=borde_grueso if col == 1 else thin,
                                         right=borde_grueso if col == num_columnas else thin)
                elif row == fin_fila_tabla:
                    cell.border = Border(top=thin, bottom=borde_grueso,
                                         left=borde_grueso if col == 1 else thin,
                                         right=borde_grueso if col == num_columnas else thin)
                elif col == 1:
                    cell.border = Border(left=borde_grueso, top=thin, bottom=thin, right=thin)
                elif col == num_columnas:
                    cell.border = Border(right=borde_grueso, top=thin, bottom=thin, left=thin)

                cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in range(1, num_columnas + 1):
            celda = ws.cell(row=inicio_fila_tabla, column=col)
            if celda.value == "Promedio(dBuV/m)":
                celda.value = "Promedio\n(dBuV/m)"
                ws.column_dimensions[get_column_letter(col)].width = 15
            elif celda.value == "Ancho de Banda (KHz)":
                celda.value = "Ancho de Banda\n(KHz)"
            elif celda.value == "Medición Manual":
                celda.value = "Medición Manual\nAB(KHz) o NIVEL\n(dBµV/m)"
                ws.column_dimensions[get_column_letter(col)].width = 23
            celda.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

        # Ajustar anchos del resto
        for i in range(1, num_columnas + 1):
            if ws.column_dimensions[get_column_letter(i)].width in [15, 23]:
                continue
            col_letter = get_column_letter(i)
            max_length = 0
            for cell in ws[col_letter]:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col_letter].width = max_length + 2

        ws.row_dimensions[inicio_fila_tabla].height = 45

    def colorear_celdas_por_valor(self, ws, fila_inicio, fila_fin, columna_inicio, columna_fin):
        """Colorea las celdas con criterios específicos por columna"""
        # Colores para el rango principal (días 1-31)
        verde = PatternFill(start_color="FFCDFECE", end_color="FFCDFECE", fill_type="solid")
        amarillo = PatternFill(start_color="FFFFFE9F", end_color="FFFFFE9F", fill_type="solid")
        rosa = PatternFill(start_color="FFFD9BCB", end_color="FFFD9BCB", fill_type="solid")
        
        # Colores especiales para columnas específicas
        rojo_ah = PatternFill(start_color="FFFF6666", end_color="FFFF6666", fill_type="solid")
        rojo_ai = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")

        # Columnas fijas AH y AI
        col_ah = 34
        col_ai = 35
        
        # 1. Colorear rango principal (días 1-31)
        for fila in range(fila_inicio, fila_fin + 1):
            for col in range(columna_inicio, columna_fin + 1):
                celda = ws.cell(row=fila, column=col)  
                try:
                    if celda.value is not None:
                        valor_str = str(celda.value).replace(',', '.').strip()
                        if valor_str.replace('.', '', 1).isdigit():
                            valor = float(valor_str)
                            
                            if 0 <= valor <= 43:
                                celda.fill = rosa
                            elif 43 < valor < 54:
                                celda.fill = amarillo
                            elif valor >= 54:
                                celda.fill = verde
                except (ValueError, TypeError):
                    pass
        
        # 2. Colorear columnas especiales
        for fila in range(fila_inicio, fila_fin + 1):
            # COLUMNA AH (34) - Rojo si > 60
            celda_ah = ws.cell(row=fila, column=col_ah)
            try:
                if celda_ah.value is not None:
                    valor_str = str(celda_ah.value).replace(',', '.').strip()
                    if valor_str.replace('.', '', 1).isdigit():
                        valor_ah = float(valor_str)
                        if valor_ah > 60:
                            celda_ah.fill = rojo_ah
            except (ValueError, TypeError):
                pass
            
            # COLUMNA AI (35) - Rojo si > 200
            celda_ai = ws.cell(row=fila, column=col_ai)
            try:
                if celda_ai.value is not None:
                    valor_str = str(celda_ai.value).replace(',', '.').strip()
                    if valor_str.replace('.', '', 1).isdigit():
                        valor_ai = float(valor_str)
                        if valor_ai > 200:
                            celda_ai.fill = rojo_ai
            except (ValueError, TypeError):
                pass

    def encontrar_columnas_numericas(self, ws, fila_encabezados):
        """Encuentra columnas con valores numéricos (días 1-31)"""
        columnas_numericas = []
        
        for col in range(1, ws.max_column + 1):
            celda = ws.cell(row=fila_encabezados, column=col)
            if celda.value and str(celda.value).isdigit():
                try:
                    dia = int(celda.value)
                    if 1 <= dia <= 31:
                        columnas_numericas.append(col)
                except ValueError:
                    pass
        
        if columnas_numericas:
            return min(columnas_numericas), max(columnas_numericas)
        else:
            return 3, 33

    def obtener_base(self, nombre_archivo):
        return nombre_archivo.split("_")[0].lower().strip()

    def process_files(self):
        try:
            ruta_fm = self.fm_path.text()
            ruta_tv = self.tv_path.text()
            ruta_salida = self.output_path.text()
            ruta_imagenes = self.img_path.text()
            
            os.makedirs(ruta_salida, exist_ok=True)
            fecha_actual = datetime.today().strftime("%d/%m/%Y")

            self.log_text.append("Leyendo archivos CSV...")
            
            # Leer archivos
            archivos_fm = {self.obtener_base(f): os.path.join(ruta_fm, f) for f in os.listdir(ruta_fm) if f.endswith(".csv")}
            archivos_tv = {self.obtener_base(f): os.path.join(ruta_tv, f) for f in os.listdir(ruta_tv) if f.endswith(".csv")}

            nombres_bases = set(archivos_fm.keys()).union(archivos_tv.keys())

            for base in nombres_bases:
                self.log_text.append(f"Procesando base: {base}")
                
                wb = Workbook()
                ws = wb.active
                ws.title = "Resumen"
                fila_actual = 1

                for tipo, archivos, titulo in [
                    ("FM", archivos_fm, "FORMULARIO DE CONTROL MENSUAL DE FM"),
                    ("TV", archivos_tv, "FORMULARIO DE CONTROL MENSUAL DE TV")
                ]:
                    if base not in archivos:
                        continue

                    ruta_archivo = archivos[base]
                    df = pd.read_csv(ruta_archivo, encoding="unicode_escape")
                    
                    if tipo == "TV":
                        df = df.rename(columns={"Nombre de la estación": "ESTACION"})
                        df = df[df.columns[:5]]
                    else:
                        df = df[df.columns[:9]]

                    df["ESTACION"] = df["ESTACION"].astype(str).str.strip()
                    df = df[df["ESTACION"].notna() & (df["ESTACION"] != "")]
                    df["Tiempo"] = pd.to_datetime(df["Tiempo"], format="%d/%m/%Y  %H:%M:%S,%f", errors='coerce').dt.date
                    
                    # Encontrar mes objetivo
                    mes_objetivo = df["Tiempo"].dropna().apply(lambda x: x.month).value_counts().idxmax()
                    nombre_mes_es = ["enero", "febrero", "marzo", "abril", "mayo", "junio", 
                                   "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"][mes_objetivo - 1]
                    
                    df = df[df["Tiempo"].apply(lambda x: x.month == mes_objetivo)]
                    df["DIA"] = pd.to_datetime(df["Tiempo"]).dt.day
                    df["Frecuencia (MHz)"] = df["Frecuencia (Hz)"] / 1_000_000
                    
                    # Convertir Level a numérico
                    df["Level (dBµV/m)"] = pd.to_numeric(df["Level (dBµV/m)"].astype(str).str.replace(",", ".", regex=False), errors='coerce')

                    # Agrupar y pivotar
                    agrupado = df.groupby(["ESTACION", "Frecuencia (MHz)", "DIA"])["Level (dBµV/m)"].mean().reset_index()
                    pivot = agrupado.pivot(index=["ESTACION", "Frecuencia (MHz)"], columns="DIA", values="Level (dBµV/m)")
                    
                    # Todos los días del mes
                    todos_los_dias = list(range(1, 32))
                    pivot = pivot.reindex(columns=todos_los_dias, fill_value=0)
                    pivot[todos_los_dias] = pivot[todos_los_dias].replace(0, "-")
                    
                    # Calcular promedio
                    pivot_numeric = pivot[todos_los_dias].replace("-", pd.NA).apply(pd.to_numeric, errors="coerce")
                    pivot["Promedio(dBuV/m)"] = pivot_numeric.mean(axis=1, skipna=True).round(2)

                    if tipo == "FM":
                        ancho_banda = df.groupby(["ESTACION", "Frecuencia (MHz)"])["Bandwidth (Hz)"].mean().reset_index()
                        ancho_banda["Ancho de Banda (KHz)"] = (ancho_banda["Bandwidth (Hz)"] / 1000).round(2)
                        pivot = pivot.reset_index().merge(ancho_banda.drop(columns=["Bandwidth (Hz)"]), on=["ESTACION", "Frecuencia (MHz)"], how="left")
                    else:
                        pivot = pivot.reset_index()

                    pivot["Medición Manual"] = ""
                    pivot["OBSERVACIONES"] = ""
                    pivot = pivot.sort_values(by="Frecuencia (MHz)")

                    # Redondear valores numéricos
                    for col in pivot.select_dtypes(include="number").columns:
                        pivot[col] = pivot[col].round(2)

                    # Escribir en Excel
                    for i, row in enumerate(dataframe_to_rows(pivot, index=False, header=True)):
                        for j, val in enumerate(row, start=1):
                            ws.cell(row=fila_actual + i, column=j, value=val)

                    # Encabezado
                    if tipo == "FM":
                        encabezado = [
                            "INFORME DE CONTROL TÉCNICO",
                            "No. IT-CZ06-R-2025-00XX",
                            "AGENCIA DE REGULACIÓN Y CONTROL DE LAS TELECOMUNICACIONES",
                            "COORDINACIÓN ZONAL 6",
                            "ESTACIÓN DE COMPROBACIÓN TÉCNICA",
                            titulo,
                            "CIUDAD:" + ("tambo" if base == "cañar" else base.upper()),
                            f"PERIODO: {nombre_mes_es.upper()}",
                            f"FECHA PRESENTACIÓN: {fecha_actual}"
                        ]
                    else:
                        encabezado = [
                            "AGENCIA DE REGULACIÓN Y CONTROL DE LAS TELECOMUNICACIONES",
                            "COORDINACIÓN ZONAL 6",
                            "ESTACIÓN DE COMPROBACIÓN TÉCNICA",
                            titulo,
                            "CIUDAD:" + ("tambo" if base == "cañar" else base.upper()),
                            f"PERIODO: {nombre_mes_es.upper()}",
                            f"FECHA PRESENTACIÓN: {fecha_actual}"
                        ]

                    self.formatear_hoja(ws, fila_actual, encabezado)

                    # Encontrar y colorear celdas
                    fila_encabezados_columnas = fila_actual + len(encabezado)
                    col_inicio_numeros, col_fin_numeros = self.encontrar_columnas_numericas(ws, fila_encabezados_columnas)
                    fila_inicio_datos = fila_encabezados_columnas + 1
                    fila_fin_datos = fila_inicio_datos + len(pivot) - 1

                    self.colorear_celdas_por_valor(ws, fila_inicio_datos, fila_fin_datos, col_inicio_numeros, col_fin_numeros)

                    fila_actual = ws.max_row + 3

                # Guardar archivo
                nombre_salida = f"{base}_ReporteUnificado.xlsx"
                wb.save(os.path.join(ruta_salida, nombre_salida))
                self.log_text.append(f"✅ Archivo generado: {nombre_salida}")

            self.log_text.append("✅ Todos los archivos procesados exitosamente!")
            self.status_bar.showMessage("Proceso completado")
            
        except Exception as e:
            error_msg = f"Error: {str(e)}"
            self.log_text.append(error_msg)
            self.status_bar.showMessage("Error en el procesamiento")
            QMessageBox.critical(self, "Error", error_msg)

def main():
    app = QApplication(sys.argv)
    window = DatabaseApp()
    sys.exit(app.exec_())

if __name__ == "__main__":
    main()