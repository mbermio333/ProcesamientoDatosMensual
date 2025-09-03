# -*- coding: utf-8 -*-
import os
import json
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image as XLImage

# ------------------ CONFIGURACIÓN GENERAL ------------------
# Cargar configuración desde archivo
CONFIG_FILE = "config.json"

def cargar_configuracion():
    """Cargar configuración desde archivo JSON"""
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r') as f:
                config = json.load(f)
                return config
        except:
            pass
    
    # Valores por defecto si no hay archivo de configuración
    return {
        "fm_path": "MedicionesFmCSV",
        "tv_path": "MedicionesTvCSV", 
        "output_path": "ReportesUnificados"
    }

# Cargar configuración al inicio
config = cargar_configuracion()
ruta_fm = config.get("fm_path", "MedicionesFmCSV")
ruta_tv = config.get("tv_path", "MedicionesTvCSV")
ruta_salida = config.get("output_path", "ReportesUnificados")
ruta_imagenes = "Img"
fecha_actual = datetime.now().strftime("%d/%m/%Y")

# ------------------ FUNCIONES AUXILIARES ------------------

def inicializar_directorios():
    """Crear los directorios necesarios si no existen"""
    os.makedirs(ruta_salida, exist_ok=True)
    os.makedirs(ruta_fm, exist_ok=True)
    os.makedirs(ruta_tv, exist_ok=True)
    os.makedirs(ruta_imagenes, exist_ok=True)

def combinar_observaciones_solo_en_tv(ws, fila_inicio_tabla, fila_fin_tabla, num_columnas, es_tv=True):
    """
    Combina horizontalmente la última y penúltima columna para cada fila SOLO si es TV.
    """
    if es_tv:
        for fila in range(fila_inicio_tabla + 1, fila_fin_tabla + 1):
            ws.merge_cells(start_row=fila, start_column=num_columnas - 1, end_row=fila, end_column=num_columnas)
        
        # También combinar el encabezado (la celda con el nombre del campo)
        ws.merge_cells(start_row=fila_inicio_tabla, start_column=num_columnas - 1, end_row=fila_inicio_tabla, end_column=num_columnas)
        celda = ws.cell(row=fila_inicio_tabla, column=num_columnas - 1)
        celda.value = "OBSERVACIONES"
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row=fila_inicio_tabla, column=num_columnas).value = None

def formatear_hoja(ws, fila_inicio, encabezado_lineas=None):
    """Formatea una hoja de cálculo con bordes y estilos"""
    if encabezado_lineas:
        num_columnas = ws.max_column
        ws.insert_rows(fila_inicio, amount=len(encabezado_lineas))

        for i, texto in enumerate(encabezado_lineas):
            fila = fila_inicio + i
            celda = ws.cell(row=fila, column=1, value=texto)
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=num_columnas)
            celda.alignment = Alignment(horizontal="center", vertical="center")
            celda.font = Font(bold=True, size=12)
    
    # Bordes
    inicio_fila_tabla = fila_inicio + (len(encabezado_lineas) if encabezado_lineas else 0)
    fin_fila_tabla = ws.max_row
    thin = Side(border_style="thin")
    borde_grueso = Side(border_style="medium")

    for row in range(inicio_fila_tabla, fin_fila_tabla + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if row == inicio_fila_tabla:
                cell.border = Border(top=borde_grueso, bottom=thin,
                                     left=borde_grueso if col == 1 else thin,
                                     right=borde_grueso if col == ws.max_column else thin)
            elif row == fin_fila_tabla:
                cell.border = Border(top=thin, bottom=borde_grueso,
                                     left=borde_grueso if col == 1 else thin,
                                     right=borde_grueso if col == ws.max_column else thin)
            elif col == 1:
                cell.border = Border(left=borde_grueso, top=thin, bottom=thin, right=thin)
            elif col == ws.max_column:
                cell.border = Border(right=borde_grueso, top=thin, bottom=thin, left=thin)

            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Ajustar anchos de columnas
    for col in range(1, ws.max_column + 1):
        celda = ws.cell(row=inicio_fila_tabla, column=col)
        if celda.value == "Promedio(dBuV/m)":
            celda.value = "Promedio\n(dBuV/m)"
            ws.column_dimensions[get_column_letter(col)].width = 15
        elif celda.value == "Ancho de Banda (KHz)":
            celda.value = "Ancho de Banda\n(KHz)"
        elif celda.value == "Medición Manual":
            celda.value = "Medición Manual\nAB(KHz) o NIVEL\n(dBµV/m)"
            ws.column_dimensions[get_column_letter(col)].width = 23
        celda.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    # Ajustar anchos del resto
    for i in range(1, ws.max_column + 1):
        if ws.column_dimensions[get_column_letter(i)].width in [15, 23]:
            continue
        col_letter = get_column_letter(i)
        max_length = 0
        for cell in ws[col_letter]:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = max_length + 2

    if encabezado_lineas:
        ws.row_dimensions[inicio_fila_tabla].height = 45

def insertar_imagenes(ws, fila_destino, tipo):
    logo_izquierda = os.path.join(ruta_imagenes, "ARCOTEL.png")
    logo_derecha = os.path.join(ruta_imagenes, "nEcuador.png")

    if os.path.exists(logo_izquierda):
        img_left = XLImage(logo_izquierda)
        img_left.width = 500
        img_left.height = 100
        ws.add_image(img_left, f"A{fila_destino}")

    if os.path.exists(logo_derecha):
        img_right = XLImage(logo_derecha)
        img_right.width = 260
        img_right.height = 120
        ws.add_image(img_right, f"AJ{fila_destino}")

# ------------------ FUNCIÓN PARA COLOREAR CELDAS ------------------

def colorear_celdas_por_valor(ws, fila_inicio, fila_fin, tipo, col_names):
    from openpyxl.styles import PatternFill

    rojo = PatternFill(start_color="FFFC4A2C", end_color="FFFC4A2C", fill_type="solid")
    amarillo = PatternFill(start_color="FFFFFE9F", end_color="FFFFFE9F", fill_type="solid")
    verde = PatternFill(start_color="FFCDFECE", end_color="FFCDFECE", fill_type="solid")
    rosa = PatternFill(start_color="FFFD9BCB", end_color="FFFD9BCB", fill_type="solid")

    encabezados = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=fila_inicio - 1, column=col).value
        if isinstance(val, str):
            val_normalizado = val.strip().replace('\n', ' ')
            encabezados[val_normalizado] = col
        elif isinstance(val, (int, float)):
            encabezados[str(val)] = col

    def obtener_banda(freq):
        if (54 <= freq <= 72) or (76 <= freq <= 88):
            return 'I'
        elif 174 <= freq <= 216:
            return 'III'
        elif (470 <= freq <= 488) or (512 <= freq <= 608):
            return 'IV'
        elif 614 <= freq <= 698:
            return 'V'
        else:
            return None

    def get_umbral_color(banda, valor):
        if banda == 'I':
            if valor < 47:
                return rosa
            elif 47 <= valor < 68:
                return amarillo
            else:
                return verde
        elif banda == 'III':
            if valor < 56:
                return rosa
            elif 56 <= valor < 71:
                return amarillo
            else:
                return verde
        elif banda in ['IV', 'V']:
            if valor < 64:
                return rosa
            elif 64 <= valor < 74:
                return amarillo
            else:
                return verde
        else:
            return None

    for fila in range(fila_inicio, fila_fin + 1):
        frecuencia_col = None
        for key in encabezados:
            if "frecuencia" in key.lower():
                frecuencia_col = encabezados[key]
                break

        frecuencia = ws.cell(row=fila, column=frecuencia_col).value
        try:
            frecuencia = float(frecuencia)
            banda = obtener_banda(frecuencia)
        except:
            banda = None

        for nombre_col in col_names:
            nombre_normalizado = str(nombre_col).strip().replace('\n', ' ')
            if nombre_normalizado not in encabezados:
                continue

            col = encabezados[nombre_normalizado]
            celda = ws.cell(row=fila, column=col)
            valor = celda.value

            try:
                valor = float(valor)
            except (ValueError, TypeError):
                continue

            if tipo == "FM":
                if nombre_normalizado == "Promedio (dBuV/m)":
                    if 0 <= valor <= 30:
                        celda.fill = rojo
                    elif 30 < valor < 54:
                        celda.fill = amarillo
                    elif valor >= 54:
                        celda.fill = verde
                elif nombre_normalizado == "Ancho de Banda (KHz)":
                    if valor <= 220:
                        celda.fill = verde
                    elif valor > 220:
                        celda.fill = rojo
                         # Escribir mensaje en la columna "OBSERVACIONES"
                        col_obs = encabezados.get("OBSERVACIONES")
                        if col_obs:
                            celda_obs = ws.cell(row=fila, column=col_obs)
                            mensaje = "Opera con ancho de banda mayor\na lo autorizado\n(medición automática)"
                            celda_obs.value = mensaje
                            celda_obs.alignment = Alignment(wrap_text=True)
                elif nombre_normalizado in [str(d) for d in range(1, 32)]:
                    if 0 <= valor <= 30:
                        celda.fill = rosa
                    elif 30 < valor < 54:
                        celda.fill = amarillo
                    elif valor >= 54:
                        celda.fill = verde

            elif tipo == "TV":
                if banda:
                    if nombre_normalizado == "Promedio (dBuV/m)" or nombre_normalizado in [str(d) for d in range(1, 32)]:
                        frecuencia_col = encabezados.get("Frecuencia (MHz)")
                        observaciones_col = encabezados.get("OBSERVACIONES")
                        # Dentro de if nombre_normalizado == "Promedio (dBuV/m)":
                        frecuencia = ws.cell(row=fila, column=frecuencia_col).value if frecuencia_col else None

                        if frecuencia:
                            banda = obtener_banda(frecuencia)
                            color = get_umbral_color(banda, valor)

                            if color:
                                celda.fill = color
                                # Agrega mensaje si es rosa (menor al mínimo por banda)
                                if color == rosa and observaciones_col:
                                    obs_cell = ws.cell(row=fila, column=observaciones_col)
                                    existing = obs_cell.value or ""
                                    mensaje = "Niveles por debajo del borde del área de cobertura principal y secundaria"
                                    
                                    if mensaje not in existing:
                                        nuevo_texto = (existing + "\n" + mensaje).strip()
                                        obs_cell.value = nuevo_texto
                                        obs_cell.alignment = Alignment(wrap_text=True, vertical="center")
                                        
                                        # Estimar número de líneas para ajustar alto de fila
                                        num_lineas = nuevo_texto.count("\n") + 1
                                        altura = num_lineas * 15  # Puedes ajustar este valor según cómo se vea
                                        ws.row_dimensions[fila].height = altura

def encontrar_columnas_numericas(ws, fila_encabezados):
    """
    Encuentra automáticamente las columnas que contienen valores numéricos
    basándose en los encabezados de días (1 al 31)
    """
    columnas_numericas = []
    
    for col in range(1, ws.max_column + 1):
        celda = ws.cell(row=fila_encabezados, column=col)
        if celda.value and str(celda.value).isdigit():
            try:
                dia = int(celda.value)
                if 1 <= dia <= 31:
                    columnas_numericas.append(col)
            except ValueError:
                pass
    
    if columnas_numericas:
        return min(columnas_numericas), max(columnas_numericas)
    else:
        # Valores por defecto si no encuentra días
        return 3, 33

def obtener_base(nombre_archivo):
    return nombre_archivo.split("_")[0].lower().strip()

# ------------------ FUNCIONES PARA CREAR HOJAS ADICIONALES ------------------

def crear_hoja_manual(wb, nombre_hoja):
    """Crea una hoja en blanco con el nombre especificado"""
    if nombre_hoja not in wb.sheetnames:
        ws = wb.create_sheet(nombre_hoja)
        # Agregar encabezados básicos
        if "FM" in nombre_hoja:
            encabezados = ["ESTACION", "Frecuencia (MHz)", "Medición Manual AB(KHz) o NIVEL (dBµV/m)", "OBSERVACIONES"]
        else:
            encabezados = ["ESTACION", "Frecuencia (MHz)", "Medición Manual AB(KHz) o NIVEL (dBµV/m)", "OBSERVACIONES"]
        
        for col, encabezado in enumerate(encabezados, 1):
            ws.cell(row=1, column=col, value=encabezado)
        
        # Aplicar formato básico
        for col in range(1, len(encabezados) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
            celda = ws.cell(row=1, column=col)
            celda.font = Font(bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center")
        
        # Aplicar bordes a la fila de encabezados
        thin = Side(border_style="thin")
        for col in range(1, len(encabezados) + 1):
            ws.cell(row=1, column=col).border = Border(top=thin, bottom=thin, left=thin, right=thin)

def crear_hoja_observaciones(wb, datos_fm, datos_tv):
    """Crea la hoja de observaciones con los datos de FM y TV"""
    if "Observaciones" in wb.sheetnames:
        ws_obs = wb["Observaciones"]
    else:
        ws_obs = wb.create_sheet("Observaciones")
    
    # Limpiar hoja existente
    ws_obs.delete_rows(1, ws_obs.max_row)
    
    fila_actual = 1
    
    # Agregar datos de FM
    if datos_fm is not None and not datos_fm.empty:
        # Encabezado para FM
        ws_obs.cell(row=fila_actual, column=1, value="TABLA FM")
        ws_obs.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=6)
        celda = ws_obs.cell(row=fila_actual, column=1)
        celda.font = Font(bold=True, size=14)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        fila_actual += 1
        
        # Encabezados de columnas para FM
        encabezados_fm = ["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", "Ancho de Banda (KHz)", 
                         "Medición Manual AB(KHz) o NIVEL (dBµV/m)", "OBSERVACIONES"]
        
        for col, encabezado in enumerate(encabezados_fm, 1):
            ws_obs.cell(row=fila_actual, column=col, value=encabezado)
            celda = ws_obs.cell(row=fila_actual, column=col)
            celda.font = Font(bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        fila_actual += 1
        
        # Datos de FM
        for _, row in datos_fm.iterrows():
            ws_obs.cell(row=fila_actual, column=1, value=row["ESTACION"])
            ws_obs.cell(row=fila_actual, column=2, value=row["Frecuencia (MHz)"])
            ws_obs.cell(row=fila_actual, column=3, value=row["Promedio(dBuV/m)"])
            ws_obs.cell(row=fila_actual, column=4, value=row["Ancho de Banda (KHz)"])
            ws_obs.cell(row=fila_actual, column=5, value=row["Medición Manual"])
            ws_obs.cell(row=fila_actual, column=6, value=row["OBSERVACIONES"])
            fila_actual += 1
        
        fila_actual += 2  # Espacio de 2 filas entre tablas
    
    # Agregar datos de TV
    if datos_tv is not None and not datos_tv.empty:
        # Encabezado para TV
        ws_obs.cell(row=fila_actual, column=1, value="TABLA TV")
        ws_obs.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=5)
        celda = ws_obs.cell(row=fila_actual, column=1)
        celda.font = Font(bold=True, size=14)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        fila_actual += 1
        
        # Encabezados de columnas para TV
        encabezados_tv = ["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", 
                         "Medición Manual AB(KHz) o NIVEL (dBµV/m)", "OBSERVACIONES"]
        
        for col, encabezado in enumerate(encabezados_tv, 1):
            ws_obs.cell(row=fila_actual, column=col, value=encabezado)
            celda = ws_obs.cell(row=fila_actual, column=col)
            celda.font = Font(bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        fila_actual += 1
        
        # Datos de TV
        for _, row in datos_tv.iterrows():
            ws_obs.cell(row=fila_actual, column=1, value=row["ESTACION"])
            ws_obs.cell(row=fila_actual, column=2, value=row["Frecuencia (MHz)"])
            ws_obs.cell(row=fila_actual, column=3, value=row["Promedio(dBuV/m)"])
            ws_obs.cell(row=fila_actual, column=4, value=row["Medición Manual"])
            ws_obs.cell(row=fila_actual, column=5, value=row["OBSERVACIONES"])
            fila_actual += 1
    
    # Aplicar formato a toda la hoja
    formatear_hoja(ws_obs, 1, None)
    
    # Colorear celdas según valores
    if datos_fm is not None and not datos_fm.empty:
        fila_inicio_fm = 3
        fila_fin_fm = fila_inicio_fm + len(datos_fm) - 1
        columnas_colorear_fm = ["Promedio (dBuV/m)", "Ancho de Banda (KHz)"]
        colorear_celdas_por_valor(ws_obs, fila_inicio_fm, fila_fin_fm, "FM", columnas_colorear_fm)
    
    if datos_tv is not None and not datos_tv.empty:
        fila_inicio_tv = fila_inicio_fm + len(datos_fm) + 4 if datos_fm is not None and not datos_fm.empty else 3
        fila_fin_tv = fila_inicio_tv + len(datos_tv) - 1
        columnas_colorear_tv = ["Promedio (dBuV/m)"]
        colorear_celdas_por_valor(ws_obs, fila_inicio_tv, fila_fin_tv, "TV", columnas_colorear_tv)

# ------------------ FUNCIÓN PRINCIPAL DE PROCESAMIENTO ------------------

def procesar_datos(callback_progreso=None, callback_log=None):
    """
    Función principal que procesa todos los datos
    """
    # Inicializar directorios
    inicializar_directorios()
    
    if callback_log:
        callback_log("Iniciando procesamiento de datos...")
        callback_log(f"Ruta FM: {ruta_fm}")
        callback_log(f"Ruta TV: {ruta_tv}")
        callback_log(f"Ruta salida: {ruta_salida}")
    
    # Obtener listas de archivos
    try:
        archivos_fm = {obtener_base(f): os.path.join(ruta_fm, f) for f in os.listdir(ruta_fm) if f.endswith(".csv")}
        archivos_tv = {obtener_base(f): os.path.join(ruta_tv, f) for f in os.listdir(ruta_tv) if f.endswith(".csv")}
        
        if callback_log:
            callback_log(f"Encontrados {len(archivos_fm)} archivos FM y {len(archivos_tv)} archivos TV")
    except Exception as e:
        if callback_log:
            callback_log(f"Error al leer archivos: {str(e)}")
        return False

    nombres_bases = set(archivos_fm.keys()).union(archivos_tv.keys())
    total_bases = len(nombres_bases)
    
    if callback_log:
        callback_log(f"Procesando {total_bases} bases de datos")
    
    # Procesar cada base
    for i, base in enumerate(nombres_bases):
        if callback_progreso:
            # Calcular progreso (0-100)
            progreso = int((i / total_bases) * 100)
            callback_progreso(progreso)
            
        if callback_log:
            callback_log(f"Procesando base: {base}")
        
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Resumen"
            fila_actual = 1
            
            # Variables para almacenar datos para la hoja de observaciones
            datos_fm = None
            datos_tv = None

            for tipo, archivos, titulo in [
                ("FM", archivos_fm, "FORMULARIO DE CONTROL MENSUAL DE FM"),
                ("TV", archivos_tv, "FORMULARIO DE CONTROL MENSUAL DE TV")
            ]:
                if base not in archivos:
                    continue

                ruta_archivo = archivos[base]
                df = pd.read_csv(ruta_archivo, encoding="unicode_escape")
                df = df.rename(columns={"Nombre de la estación": "ESTACION"}) if tipo == "TV" else df
                df = df[df.columns[:9]] if tipo == "FM" else df[df.columns[:5]]

                df["ESTACION"] = df["ESTACION"].astype(str).str.strip()
                df = df[df["ESTACION"].notna() & (df["ESTACION"] != "")]
                df["Tiempo"] = pd.to_datetime(df["Tiempo"], format="%d/%m/%Y  %H:%M:%S,%f", errors='coerce').dt.date
                mes_objetivo = df["Tiempo"].dropna().apply(lambda x: x.month).value_counts().idxmax()

                nombre_mes_es = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"][mes_objetivo - 1]
                df = df[df["Tiempo"].apply(lambda x: x.month == mes_objetivo)]
                df["DIA"] = pd.to_datetime(df["Tiempo"]).dt.day

                df["Frecuencia (MHz)"] = df["Frecuencia (Hz)"] / 1_000_000
                df["Level (dBµV/m)"] = df["Level (dBµV/m)"].astype(str).str.replace(",", ".", regex=False).astype(float)

                agrupado = df.groupby(["ESTACION", "Frecuencia (MHz)", "DIA"])["Level (dBµV/m)"].mean().reset_index()
                pivot = agrupado.pivot(index=["ESTACION", "Frecuencia (MHz)"], columns="DIA", values="Level (dBµV/m)")
                todos_los_dias = list(range(1, 32))
                pivot = pivot.reindex(columns=todos_los_dias, fill_value=0)

                pivot[todos_los_dias] = pivot[todos_los_dias].replace(0, "-")
                pivot_numeric = pivot[todos_los_dias].replace("-", pd.NA).apply(pd.to_numeric, errors="coerce")
                pivot["Promedio(dBuV/m)"] = pivot_numeric.mean(axis=1, skipna=True).round(2)

                if tipo == "FM":
                    ancho_banda = (
                        df.groupby(["ESTACION", "Frecuencia (MHz)"])["Bandwidth (Hz)"].mean().reset_index()
                    )
                    ancho_banda["Ancho de Banda (KHz)"] = (ancho_banda["Bandwidth (Hz)"] / 1000).round(2)
                    pivot = pivot.reset_index().merge(ancho_banda.drop(columns=["Bandwidth (Hz)"]), on=["ESTACION", "Frecuencia (MHz)"], how="left")
                    # Guardar datos para hoja de observaciones
                    datos_fm = pivot[["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", "Ancho de Banda (KHz)", "Medición Manual", "OBSERVACIONES"]].copy()
                else:
                    pivot = pivot.reset_index()
                    # Guardar datos para hoja de observaciones
                    datos_tv = pivot[["ESTACION", "Frecuencia (MHz)", "Promedio(dBuV/m)", "Medición Manual", "OBSERVACIONES"]].copy()

                pivot["Medición Manual"] = ""
                pivot["OBSERVACIONES"] = ""
                pivot = pivot.sort_values(by="Frecuencia (MHz)")
                pivot = pivot.reset_index(drop=True)

                for col in pivot.select_dtypes(include="number").columns:
                    pivot[col] = pivot[col].round(2)

                for i, row in enumerate(dataframe_to_rows(pivot, index=False, header=True)):
                    for j, val in enumerate(row, start=1):
                        ws.cell(row=fila_actual + i, column=j, value=val)

                encabezado_fm = [
                    "INFORME DE CONTROL TÉCNICO",
                    "No. IT-CZ06-R-2025-00XX",
                    "AGENCIA DE REGULACIÓN Y CONTROL DE LAS TELECOMUNICACIONES",
                    "COORDINACIÓN ZONAL 6",
                    "ESTACIÓN DE COMPROBACIÓN TÉCNICA",
                    titulo,
                    "CIUDAD:" + ("tambo" if base == "cañar" else base.upper()),
                    f"PERIODO: {nombre_mes_es.upper()}",
                    f"FECHA PRESENTACIÓN: {fecha_actual}"
                ] if tipo == "FM" else [
                    "AGENCIA DE REGULACIÓN Y CONTROL DE LAS TELECOMUNICACIONES",
                    "COORDINACIÓN ZONAL 6",
                    "ESTACIÓN DE COMPROBACIÓN TÉCNICA",
                    titulo,
                    "CIUDAD:" + ("tambo" if base == "cañar" else base.upper()),
                    f"PERIODO: {nombre_mes_es.upper()}",
                    f"FECHA PRESENTACIÓN: {fecha_actual}"
                ]
                
                if tipo == "FM":
                    insertar_imagenes(ws, fila_actual + 4, tipo="FM")
                elif tipo == "TV":
                    insertar_imagenes(ws, fila_actual + 2, tipo="TV")

                formatear_hoja(ws, fila_actual, encabezado_fm)

                # Colorear celdas
                fila_encabezados_columnas = fila_actual + len(encabezado_fm)
                fila_inicio_datos = fila_encabezados_columnas + 1
                fila_fin_datos = fila_inicio_datos + len(pivot) - 1

                columnas_colorear = ["Promedio (dBuV/m)", "Ancho de Banda\n(KHz)", *list(range(1, 32))]
                colorear_celdas_por_valor(ws, fila_inicio_datos, fila_fin_datos, tipo, columnas_colorear)

                fila_actual = ws.max_row + 3

            # Crear hojas adicionales
            crear_hoja_manual(wb, "Manual FM")
            crear_hoja_manual(wb, "Manual TV")
            crear_hoja_observaciones(wb, datos_fm, datos_tv)
            
            # Reordenar hojas
            orden_hojas = ["Resumen", "Manual FM", "Manual TV", "Observaciones"]
            for hoja in orden_hojas:
                if hoja in wb.sheetnames:
                    wb.move_sheet(hoja, -len(orden_hojas))
                    orden_hojas.remove(hoja)
            
            nombre_salida = f"{base}_ReporteUnificado.xlsx"
            
            # Combinar observaciones para TV
            for sheet in wb.worksheets:
                if "Resumen" not in sheet.title:
                    continue

                # Buscar fila del segundo encabezado (TV)
                for fila in range(1, sheet.max_row + 1):
                    val = sheet.cell(row=fila, column=1).value
                    if isinstance(val, str) and "FORMULARIO DE CONTROL MENSUAL DE TV" in val:
                        fila_encabezado_tv = fila
                        break
                else:
                    continue  # No hay TV en esta hoja

                # Buscar fila exacta de la cabecera de la tabla (donde dice "ESTACION")
                fila_tabla_tv = None
                for fila in range(fila_encabezado_tv + 1, sheet.max_row + 1):
                    if sheet.cell(row=fila, column=1).value == "ESTACION":
                        fila_tabla_tv = fila
                        break

                if fila_tabla_tv is None:
                    continue  # Seguridad por si no se encuentra

                ultima_fila_tv = sheet.max_row
                col_final = sheet.max_column
                col_penultima = col_final - 1

                # Obtener anchos actuales
                ancho_col1 = sheet.column_dimensions[get_column_letter(col_penultima)].width
                ancho_col2 = sheet.column_dimensions[get_column_letter(col_final)].width
                ancho_combinado = (ancho_col1 or 10) + (ancho_col2 or 10)

                # Combinar celdas del cuerpo de la tabla
                for fila in range(fila_tabla_tv + 1, ultima_fila_tv + 1):
                    sheet.merge_cells(start_row=fila, start_column=col_penultima, end_row=fila, end_column=col_final)

                # Combinar encabezado de tabla ("OBSERVACIONES")
                sheet.merge_cells(start_row=fila_tabla_tv, start_column=col_penultima, end_row=fila_tabla_tv, end_column=col_final)
                celda_obs = sheet.cell(row=fila_tabla_tv, column=col_penultima)
                celda_obs.value = "OBSERVACIONES"
                celda_obs.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                # Asignar nuevo ancho
                sheet.column_dimensions[get_column_letter(col_penultima)].width = ancho_combinado
                ws.sheet_view.showGridLines = False
                
            wb.save(os.path.join(ruta_salida, nombre_salida))
            
            if callback_log:
                callback_log(f"✅ Archivo generado: {nombre_salida}")
                
        except Exception as e:
            if callback_log:
                callback_log(f"❌ Error procesando base {base}: {str(e)}")
    
    if callback_progreso:
        callback_progreso(100)
        
    if callback_log:
        callback_log("Procesamiento completado")
    
    return True

# ------------------ EJECUCIÓN DIRECTA (para testing) ------------------

if __name__ == "__main__":
    # Si se ejecuta directamente, usar callbacks simples
    def mostrar_progreso(progreso):
        print(f"Progreso: {progreso}%")
    
    def mostrar_log(mensaje):
        print(mensaje)
    
    # Procesar datos
    resultado = procesar_datos(mostrar_progreso, mostrar_log)
    
    if resultado:
        print("Procesamiento completado con éxito")
    else:
        print("Ocurrieron errores durante el procesamiento")